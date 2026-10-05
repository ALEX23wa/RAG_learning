"""Structural extraction. Never execute document content."""
import ast
import csv
import hashlib
import io
import json
import re
import uuid
from pathlib import Path

SUPPORTED = {'.pdf','.docx','.doc','.xlsx','.xls','.csv','.png','.jpg','.jpeg','.webp','.bmp',
             '.py','.java','.js','.ts','.c','.cpp','.h','.cs','.go','.rs','.sql','.json','.yaml','.yml','.md','.txt','.xml','.html'}
IMAGES = {'.png','.jpg','.jpeg','.webp','.bmp'}

def cached_json(path):
    """读取可复用的 JSON 检查点。

    Args:
        path: Path：缓存 JSON 文件路径。

    Returns:
        Any | None：解码后的 JSON；文件不存在或内容损坏时返回 None。

    Notes:
        权限错误等其他 I/O 异常继续向调用方抛出。
    """
    try:return json.loads(path.read_text(encoding='utf-8'))
    except (FileNotFoundError,json.JSONDecodeError):return None

def save_json(path,data):
    """通过同目录临时文件替换写入 JSON 检查点。

    Args:
        path: Path：目标路径，父目录需已存在。
        data: Any：可被 json.dumps 序列化的数据。

    Returns:
        None：写入 UTF-8 JSON 文件。

    Notes:
        先完整写入 .tmp，再替换目标，避免解析中断留下半份 JSON。
    """
    temporary=path.with_suffix(path.suffix+'.tmp')
    temporary.write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8')
    temporary.replace(path)

def excel_header(ws):
    """跳过单元格合并标题，识别工作表前部的表头。

    Args:
        ws: openpyxl 工作表：应至少包含一行；最多检查前 20 行。

    Returns:
        tuple[int, str]：从 1 开始的表头行号，以及含单元格坐标的表头文本。

    Notes:
        优先选择至少两个非空单元格且半数以上为字符串的行；未找到时回退第一行。
    """
    for number,row in enumerate(ws.iter_rows(min_row=1,max_row=min(20,ws.max_row)),1):
        filled=[c for c in row if c.value not in (None,'')]
        if len(filled)>=2 and sum(isinstance(c.value,str) for c in filled)>=len(filled)/2:
            return number,' | '.join(f'{c.coordinate}={c.value}' for c in filled)
    row=next(ws.iter_rows())
    return 1,' | '.join(f'{c.coordinate}={c.value}' for c in row if c.value is not None)

def block(text, locator, image=None):
    """将各格式解析结果统一为结构化内容块。

    Args:
        text: Any：正文内容，转换为字符串并保留原始空白、换行和缩进。
        locator: dict：原文件物理定位，如 page、paragraph、sheet/range 或 line_start/line_end。
        image: Path | str | None：需要视觉复核的图片路径。

    Returns:
        dict：含 text、locator、image；image 为字符串路径或 None。
    """
    return {'text': str(text), 'locator': locator, 'image': str(image) if image else None}

def read_text(path):
    """依次尝试项目支持的文本编码，避免静默丢弃非法字符。

    Args:
        path: Path：待读取的文本文件。

    Returns:
        str：完整解码文本。

    Raises:
        ValueError：UTF-8-SIG、UTF-16、GB18030 均无法严格解码。
    """
    raw = path.read_bytes()
    for codec in ('utf-8-sig','utf-16','gb18030'):
        try:
            return raw.decode(codec)
        except UnicodeError:
            pass
    raise ValueError('Cannot decode text without data loss')

def ocr(image):
    """通过 Tesseract 对图片进行简体中文和英文 OCR。

    Args:
        image: Path | str：本地图片文件。

    Returns:
        str：OCR 文本，可能包含识别错误，不等同于视觉事实确认。

    Notes:
        优先使用本机 E 盘 Tesseract；单次识别超时 120 秒，缺失程序或语言包等错误会向上传递。
    """
    import pytesseract
    from PIL import Image
    binary = Path(r'E:\Program Files\tesseract\tesseract.exe')
    if binary.exists():
        pytesseract.pytesseract.tesseract_cmd = str(binary)
    with Image.open(image) as im:
        return pytesseract.image_to_string(im, lang='chi_sim+eng', timeout=120)

def pdf_quality(page, text):
    """Inspect text quality and column positions without modifying the source."""
    flags = []
    if len(re.findall(r'[\w\u4e00-\u9fff]', text)) < 80:
        flags.append('sparse_text')
    if text.count('\ufffd') > max(3, len(text) * .01):
        flags.append('garbled_text')
    lines = [line for line in text.splitlines() if line.strip()]
    if sum(bool(re.search(r'\S[ \t]{3,}\S', line)) for line in lines) >= 3:
        flags.append('aligned_columns_or_table')
    if sum(line.count('|') >= 2 or line.count('\t') >= 2 for line in lines) >= 2:
        flags.append('table_delimiters')
    positions = []
    def visitor(fragment, cm, tm, font, size):
        if fragment.strip(): positions.append((float(tm[4]), float(tm[5])))
    try:
        page.extract_text(visitor_text=visitor)
        width = float(page.mediabox.width)
        left = [y for x, y in positions if x < width * .42]
        right = [y for x, y in positions if x > width * .48]
        if len(left) >= 4 and len(right) >= 4 and any(abs(a-b)<15 for a in left for b in right):
            flags.append('multiple_text_columns')
    except (AttributeError, TypeError, ValueError, KeyError):
        flags.append('layout_probe_unavailable')
    return flags


def pdf(path, cache, settings=None):
    """Parse text/layout/auto routes, with document or physical-page overrides.

    Auto uses text quality, aligned columns and text-position signals. Flags are
    diagnostic heuristics, not a guarantee of correct reading order. Explicit
    layout mode uses MinerU even on text-rich pages; failed parsing is not hidden.
    """
    from pypdf import PdfReader, PdfWriter
    from data_process.data_md import pdf_to_markdown_mineru
    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(''):
        raise ValueError('Encrypted PDF')
    mode = getattr(settings, 'pdf_mode', 'auto')
    overrides = getattr(settings, 'pdf_overrides', {}) or {}
    relative = path.name
    if settings is not None and getattr(settings, 'corpus', None):
        try: relative = path.relative_to(settings.corpus).as_posix()
        except ValueError: pass
    override = overrides.get(relative, overrides.get(path.name, mode))
    if isinstance(override, str): mode = override
    if mode not in {'auto', 'text', 'layout'}:
        raise ValueError('PDF mode must be auto, text or layout')
    with path.open('rb') as stream:
        source_hash=hashlib.file_digest(stream,'sha256').hexdigest()
    if isinstance(override,dict) and override.get('source_sha256',source_hash)!=source_hash:
        raise ValueError('Reviewed PDF override does not match current source bytes; review again')
    provenance={'source_sha256':source_hash,'mode':mode,'override':override,'parser':'layout-signals-v2'}
    route_key = hashlib.sha256(json.dumps(provenance, sort_keys=True).encode()).hexdigest()[:16]
    cache = cache / ('pdf-' + route_key)
    cache.mkdir(parents=True, exist_ok=True)
    def read_page(number):
        saved=cached_json(cache/f'page-{number}.json')
        if isinstance(saved,dict) and saved.get('provenance')=={**provenance,'page':number}:
            blocks=saved.get('blocks')
            checksum=hashlib.sha256(json.dumps(blocks,ensure_ascii=False,sort_keys=True).encode()).hexdigest()
            if saved.get('blocks_sha256')==checksum:return blocks
        return None
    def write_page(number,blocks):
        save_json(cache/f'page-{number}.json',{'provenance':{**provenance,'page':number},'blocks':blocks,
            'blocks_sha256':hashlib.sha256(json.dumps(blocks,ensure_ascii=False,sort_keys=True).encode()).hexdigest()})
    pending=[]; page_info={}
    for index,page in enumerate(reader.pages):
        if read_page(index+1) is not None: continue
        page_mode=override.get(str(index+1),override.get('default',mode)) if isinstance(override,dict) else mode
        if page_mode not in {'auto','text','layout'}: raise ValueError('Invalid PDF page mode')
        # A review may override one physical page. Reuse other pages only when
        # source bytes, parser and effective page route are identical, with a
        # verified block checksum. Different routes never share a checkpoint.
        reused=False
        for checkpoint in sorted(cache.parent.glob(f'pdf-*/page-{index+1}.json')):
            if checkpoint.parent==cache:continue
            saved=cached_json(checkpoint)
            if not isinstance(saved,dict):continue
            prior=saved.get('provenance',{});prior_override=prior.get('override',prior.get('mode'))
            prior_mode=(prior_override.get(str(index+1),prior_override.get('default',prior.get('mode')))
                        if isinstance(prior_override,dict) else prior_override)
            if (prior.get('source_sha256')!=source_hash or prior.get('parser')!=provenance['parser']
                or prior.get('page')!=index+1 or prior_mode!=page_mode):continue
            blocks=saved.get('blocks')
            if not isinstance(blocks,list) or saved.get('blocks_sha256')!=hashlib.sha256(
                json.dumps(blocks,ensure_ascii=False,sort_keys=True).encode()).hexdigest():continue
            write_page(index+1,blocks);reused=True;break
        if reused:continue
        text=page.extract_text() or ''
        flags=pdf_quality(page,text)
        needs_layout=page_mode=='layout' or (page_mode=='auto' and any(f!='layout_probe_unavailable' for f in flags))
        page_info[index]=(flags,page_mode)
        if needs_layout: pending.append(index)
        else:
            write_page(index+1,[block(text, {'page':index+1,'method':'pdf_text',
                'quality_flags':flags,'parse_mode':page_mode,'section_path':[]})])
    for begin in range(0,len(pending),8):
        indices=pending[begin:begin+8]
        target=cache/('batch-'+hashlib.sha256(str(indices).encode()).hexdigest()[:12])
        target.mkdir(exist_ok=True)
        single=target/'input.pdf'
        writer=PdfWriter()
        for i in indices: writer.add_page(reader.pages[i])
        buffer=io.BytesIO();writer.write(buffer);input_bytes=buffer.getvalue()
        batch_provenance={**provenance,'original_pages':[i+1 for i in indices],
                          'input_sha256':hashlib.sha256(input_bytes).hexdigest()}
        manifest=cached_json(target/'provenance.json');data=None;artifact_dir=None
        if isinstance(manifest,dict) and manifest.get('input')==batch_provenance:
            artifact_dir=target/manifest.get('artifact_directory','missing')
            resolved=artifact_dir.resolve()
            if not resolved.is_relative_to(target.resolve()):raise ValueError('Unsafe MinerU cache path')
            result_path=artifact_dir/'structured_content.json'
            if result_path.is_file() and hashlib.sha256(result_path.read_bytes()).hexdigest()==manifest.get('output_sha256'):
                data=cached_json(result_path)
        if data is None:
            single.write_bytes(input_bytes)
            # A fresh output directory prevents stale files from a previous input
            # or a failed worker from being treated as this run's successful output.
            artifact_dir=target/('run-'+uuid.uuid4().hex)
            artifact_dir.mkdir()
            pdf_to_markdown_mineru(single,artifact_dir,device='auto')
            result_path=artifact_dir/'structured_content.json'
            data=json.loads(result_path.read_text(encoding='utf-8'))
        if (len(data.get('pages',[]))!=len(indices) or
            {p.get('page_idx') for p in data['pages']}!=set(range(len(indices)))):
            raise ValueError('MinerU did not return every requested page exactly once')
        save_json(target/'provenance.json',{'input':batch_provenance,
            'artifact_directory':artifact_dir.name,
            'output_sha256':hashlib.sha256((artifact_dir/'structured_content.json').read_bytes()).hexdigest()})
        for page in data['pages']:
            original=indices[page['page_idx']]; extracted=[]; heading=''
            for n,item in enumerate(page['blocks']):
                body=item.get('content','')
                if not isinstance(body,str): body=json.dumps(body,ensure_ascii=False)
                captions='\n'.join(str(c) for c in item.get('captions',[]))
                if captions: body+='\n'+captions
                picture=artifact_dir/item['image_source'] if item.get('image_source') else None
                if picture and not picture.resolve().is_relative_to(artifact_dir.resolve()):
                    raise ValueError('MinerU image path escapes its output directory')
                if body.strip() or picture:
                    if item.get('type') in {'title','heading','paragraph_title'}: heading=body.strip()
                    locator={'page':original+1,'block':n+1,'type':item.get('type','text'),'method':'mineru',
                             'quality_flags':page_info[original][0],'parse_mode':page_info[original][1],
                             'section_path':[heading] if heading else [],'content_origin':'model_extraction'}
                    if item.get('bbox') is not None: locator['bbox']=item['bbox']
                    extracted.append(block(body or '[PDF image; requires visual verification]',locator,picture))
            if not extracted: raise ValueError(f'MinerU returned no content for page {original+1}')
            write_page(original+1,extracted)
    output=[]
    for index in range(len(reader.pages)):
        existing=read_page(index+1)
        if existing is None: raise ValueError(f'PDF page {index+1} has no completed parsing checkpoint')
        output.extend(existing)
    return output


def parse(path, cache, settings=None):
    """按文件扩展名提取正文和来源定位，统一交给分块器。

    Args:
        path: Path：待解析的原文件；调用方负责限定其位于允许的资料目录。
        cache: Path：解析资产输出目录，函数会自动创建。

    Returns:
        list[dict]：每块含 text、locator、image；页/段落/表格行/源码行等定位随格式变化。

    Raises:
        ValueError：旧 .doc、无法解码的文本或不支持直接处理的文件内容；第三方解析器也可能抛出格式错误。

    Notes:
        图片和 DOCX 内嵌图保存 OCR；Excel 保留公式与缓存状态；Python 仅做 AST 解析，不执行源码。旧式 .doc 明确报转换需求。
    """
    cache.mkdir(parents=True, exist_ok=True)
    suffix = path.suffix.lower()
    if suffix == '.pdf': return pdf(path, cache, settings)
    if suffix in IMAGES:
        return [block('[OCR; may contain recognition errors]\n'+ocr(path), {'image':path.name}, path)]
    if suffix == '.docx':
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = Document(path); result=[]; pn=tn=0; heading=''; sections=[]
        for item in doc.iter_inner_content():
            if isinstance(item, Paragraph):
                pn += 1
                if item.style and item.style.name.startswith('Heading'):
                    heading=item.text
                    level=int(item.style.name.split()[-1]) if item.style.name.split()[-1].isdigit() else 1
                    sections=sections[:level-1]+[heading]
                if item.text.strip(): result.append(block(item.text, {'paragraph':pn, 'heading':heading,
                    'section_path':list(sections),'kind':'paragraph'}))
            elif isinstance(item, Table):
                tn += 1
                headers=' | '.join(c.text for c in item.rows[0].cells) if item.rows else ''
                for rown, row in enumerate(item.rows, 1):
                    result.append(block(' | '.join(c.text for c in row.cells), {'table':tn, 'row':rown,
                        'table_id':f'docx-{tn}','table_header':headers,'kind':'table_row',
                        'header_locator':{'table':tn,'row':1},'section_path':list(sections)}))
        for i, rel in enumerate(doc.part.rels.values()):
            if 'image' in rel.reltype and not rel.is_external:
                name = cache / f'docx-image-{i}{Path(rel.target_ref).suffix}'
                name.write_bytes(rel.target_part.blob)
                try: text=ocr(name)
                except Exception: text='OCR unavailable; requires visual verification'
                result.append(block('[Embedded image OCR]\n'+text, {'embedded_image':i}, name))
        return result
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        wb=load_workbook(path,read_only=True,data_only=False)
        values=load_workbook(path,read_only=True,data_only=True); result=[]
        try:
            for ws in wb:
                header_row,headers=excel_header(ws)
                units=[]
                for row in ws.iter_rows(min_row=header_row+1,max_row=min(header_row+2,ws.max_row)):
                    for c in row:
                        if isinstance(c.value,str) and re.fullmatch(r'\s*[（(][^）)]+[）)]\s*',c.value):units.append(f'{c.coordinate}={c.value}')
                for number,(row, cached) in enumerate(zip(ws.iter_rows(), values[ws.title].iter_rows()),1):
                    if not any(c.value is not None for c in row): continue
                    fields=[]
                    for c,v in zip(row,cached):
                        if c.value is not None:
                            val=f'{c.coordinate}={c.value}'
                            if c.data_type=='f': val+=f' [formula cache={v.value!r}; '+('missing' if v.value is None else 'unverified freshness')+']'
                            fields.append(val)
                    body=' | '.join(fields)
                    item=block('Sheet: '+ws.title+'\nHeaders: '+headers+'\nUnits: '+' | '.join(units)+'\n'+body,
                        {'sheet':ws.title,'range':f'A{number}:{get_column_letter(len(row))}{number}',
                         'row':number,'table_id':ws.title,'table_header':headers,'kind':'table_row',
                         'header_locator':{'sheet':ws.title,'range':f'A{header_row}:{get_column_letter(len(row))}{header_row}','row':header_row},
                         'section_path':[ws.title],'units':units})
                    item['row_text']=body;result.append(item)
        finally: wb.close(); values.close()
        return result
    if suffix == '.xls':
        import xlrd
        wb=xlrd.open_workbook(path); result=[]
        for ws in wb.sheets():
            for i in range(ws.nrows):
                result.append(block(f'Headers: {ws.row_values(0)}\n{ws.row_values(i)}',
                    {'sheet':ws.name,'row':i+1,'formula_cache':'legacy values only','table_id':ws.name,
                     'table_header':str(ws.row_values(0)),'header_locator':{'sheet':ws.name,'row':1},
                     'section_path':[ws.name],'kind':'table_row'}))
        return result
    if suffix == '.doc':
        raise ValueError('Legacy .doc requires conversion to .docx using Word or LibreOffice')
    if suffix == '.csv':
        reader=csv.reader(io.StringIO(read_text(path))); headers=next(reader, []); result=[]
        for i,row in enumerate(reader,2):
            if not any(row): continue
            if len(row)!=len(headers):raise ValueError(f'CSV row {i} has {len(row)} fields; expected {len(headers)}')
            result.append(block(' | '.join(f'{h}={v}' for h,v in zip(headers,row)),
                {'row':i,'columns':headers,'table_id':'csv','table_header':' | '.join(headers),
                 'header_locator':{'row':1,'columns':headers},'kind':'table_row'}))
        return result
    text=read_text(path); lines=text.splitlines(keepends=True); result=[]
    line_offsets=[]; char_offset=0
    for line in lines:line_offsets.append(char_offset);char_offset+=len(line)
    def text_block(start,end,extra=None):
        locator={'line_start':start,'line_end':end,'document_char_start':line_offsets[start-1],
                 'kind':'code' if suffix in {'.py','.java','.js','.ts','.c','.cpp','.h','.cs','.go','.rs','.sql'} else 'text',
                 'structure':'text_fallback'}
        locator.update(extra or {})
        return block(''.join(lines[start-1:end]),locator)
    if suffix=='.py':
        try:
            tree=ast.parse(text); previous=1
            def ast_block(start,end,extra=None):
                nodes=[n for n in ast.walk(tree) if isinstance(n,ast.stmt)
                       and start<=n.lineno<=getattr(n,'end_lineno',n.lineno)<=end]
                ends=sorted({sum(len(s) for s in lines[start-1:n.end_lineno]) for n in nodes})
                imports=[ast.get_source_segment(text,n) for n in nodes if isinstance(n,(ast.Import,ast.ImportFrom))]
                return text_block(start,end,{'structure':'python_ast','statement_ends':ends,
                    'imports':imports,**(extra or {})})
            for node in tree.body:
                end=getattr(node,'end_lineno',node.lineno)
                if isinstance(node,(ast.FunctionDef,ast.AsyncFunctionDef,ast.ClassDef)):
                    start=min([node.lineno]+[d.lineno for d in node.decorator_list])
                    if previous<start: result.append(ast_block(previous,start-1))
                    result.append(ast_block(start,end,{'symbol':node.name,'section_path':[node.name],
                        'symbol_kind':'class' if isinstance(node,ast.ClassDef) else 'function'}))
                    previous=end+1
            if previous<=len(lines): result.append(ast_block(previous,len(lines)))
            return result
        except SyntaxError: pass
    start=1; heading=''; fenced=False;sections=[];block_sections=[]
    for i,line in enumerate(lines,1):
        if line.lstrip().startswith('```'): fenced=not fenced
        if i>start and not fenced and (line.startswith('#') or (not line.strip() and i-start>=12)):
            result.append(text_block(start,i-1,{'heading':heading,'section_path':list(block_sections)}))
            start=i
        if line.startswith('#') and not fenced:
            heading=line.lstrip('# ').strip();level=len(line)-len(line.lstrip('#'))
            sections=sections[:max(0,level-1)]+[heading]
        if start==i:block_sections=list(sections)
    if start<=len(lines): result.append(text_block(start,len(lines),{'heading':heading,'section_path':list(block_sections)}))
    return result

def chunk(blocks, tokenizer, maximum=384, overlap=48):
    """Yield bounded structure chunks with exact source spans and raw text."""
    from .structure import structural_chunks
    yield from structural_chunks(blocks, tokenizer, maximum, overlap)
