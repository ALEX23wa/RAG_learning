"""Bounded table lookup/aggregation, with physical coordinates and no eval/exec."""
import csv
import io
import math
import re
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path

from .parsers import read_text

OPERATIONS = {'sum', 'mean', 'min', 'max', 'count'}
MAX_ROWS = 1_000_000
UNIT = re.compile(r'\s*[（(]([^）)]+)[）)]\s*')


def prompt_calculation(result):
    """Losslessly compress row membership for prompts, retaining the full audit result.

    An aggregate may include a million rows. Sending that list twice can crowd
    out the actual value. Inclusive ranges preserve every included/missing row;
    the returned calculation and administrative trace still retain full lists.
    """
    compact = dict(result)
    for key in ('included_rows', 'missing_rows'):
        rows = result.get(key, [])
        if len(rows) <= 32:
            continue
        ranges = []
        previous = 0
        for row in rows:
            if type(row) is not int or row <= previous:
                raise ValueError('Calculation row membership must be positive and strictly increasing')
            if ranges and row == ranges[-1][1] + 1:
                ranges[-1][1] = row
            else:
                ranges.append([row, row])
            previous = row
        compact.pop(key)
        compact[key + '_ranges'] = ranges
        compact[key + '_count'] = len(rows)
    return compact


def _serial(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError('Non-finite table value')
    return value


def _column_index(header, column):
    matches = [i for i, value in enumerate(header) if value == column]
    if len(matches) != 1:
        raise ValueError('Column must exactly match one unique table header')
    return matches[0]


@contextmanager
def _table(path, sheet=None):
    """Open originals only; each row has a physical number and formula cells."""
    path = Path(path)
    if path.suffix.lower() == '.csv':
        if sheet is not None:
            raise ValueError('CSV has no worksheets')
        reader = csv.reader(io.StringIO(read_text(path), newline=''))
        try:
            header = next(reader)
        except StopIteration:
            raise ValueError('Empty table') from None
        if not header or not any(header):
            raise ValueError('Empty table header')
        if reader.line_num != 1:
            raise ValueError('Multi-line CSV records require structured record selection')

        def records():
            yield 1, header, None
            previous = 1
            for values in reader:
                current = reader.line_num
                if current != previous + 1:
                    raise ValueError('Multi-line CSV records require structured record selection')
                previous = current
                if current > MAX_ROWS + 1:
                    raise ValueError('Table exceeds one million rows')
                if len(values) > len(header):
                    raise ValueError('CSV row has cells outside the declared header')
                yield current, values, None

        yield {'header': header, 'header_row': 1, 'sheet': None, 'records': records}
        return
    if path.suffix.lower() != '.xlsx':
        raise ValueError('Calculator supports CSV and XLSX')
    from openpyxl import load_workbook
    from .parsers import excel_header
    cached = load_workbook(path, data_only=True, read_only=True)
    formulas = None
    try:
        formulas = load_workbook(path, data_only=False, read_only=True)
        if sheet is not None and sheet not in cached.sheetnames:
            raise ValueError('Unknown worksheet')
        ws = cached[sheet] if sheet else cached.active
        header_row, _ = excel_header(ws)
        header = [str(v) if v is not None else '' for v in next(
            ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]

        def records():
            if ws.max_row > MAX_ROWS + header_row:
                raise ValueError('Table exceeds one million rows')
            for number, (values, cells) in enumerate(zip(
                    ws.iter_rows(values_only=True), formulas[ws.title].iter_rows()), 1):
                yield number, values, cells

        yield {'header': header, 'header_row': header_row, 'sheet': ws.title, 'records': records,
               'worksheet_xml': getattr(ws, '_worksheet_path', None)}
    finally:
        cached.close()
        if formulas is not None:
            formulas.close()


def _cell(values, cells, index, number):
    value = values[index] if index < len(values) else None
    if cells is not None and index < len(cells):
        cell = cells[index]
        if cell.data_type == 'f' and value is None:
            raise ValueError(f'Missing formula cache at {cell.coordinate}; recalculate in Excel first')
        if cell.data_type == 'e':
            raise ValueError(f'Excel error at {cell.coordinate}')
    return _serial(value)


def _units(header):
    return {h: m.group(1) for h in header if h and (m := re.search(r'[（(]([^）)]+)[）)]\s*$', h))}


def _unit_row(values, number, header_row):
    filled = [v for v in values if v not in (None, '')]
    return (header_row < number <= header_row + 2 and bool(filled)
            and all(isinstance(v, str) and (UNIT.fullmatch(v) or v.strip() in {'单位', 'unit', 'units'}) for v in filled))


def _accepts(value, condition):
    if isinstance(value, str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}T00:00:00', value):
        value = value[:10]
    if value in (None, ''):
        return False
    if not isinstance(condition, dict):
        return str(value) == str(condition)
    for operator, target in condition.items():
        if operator == 'prefix' and not str(value).startswith(str(target)):
            return False
        if operator == 'eq' and str(value) != str(target):
            return False
        if operator in {'gte', 'lte'}:
            try:
                a, b = float(value), float(target)
                if not math.isfinite(a) or not math.isfinite(b):
                    raise ValueError('Non-finite filter value')
            except (TypeError, ValueError):
                try:
                    a, b = date.fromisoformat(str(value)), date.fromisoformat(str(target))
                except ValueError:
                    raise ValueError('Range filters require finite numbers or ISO dates') from None
            if operator == 'gte' and a < b or operator == 'lte' and a > b:
                return False
    return True


def calculate(path, operation, column, *, sheet=None, filters=None):
    """Aggregate a unique column with explicit AND filters and included rows.

    Blank targets are excluded. Missing formula caches in selected targets or
    filter cells are errors. Unknown operators are rejected even for empty tables.
    Original units are reported without implicit conversion.
    """
    if operation not in OPERATIONS:
        raise ValueError('Unsupported operation')
    filters = {} if filters is None else filters
    if not isinstance(filters, dict):
        raise ValueError('Filters must be a mapping')
    for condition in filters.values():
        if isinstance(condition, dict) and (not condition or set(condition) - {'eq', 'prefix', 'gte', 'lte'}):
            raise ValueError('Unsupported filter operator')
    with _table(path, sheet) as table:
        header = table['header']
        index = _column_index(header, column)
        indices = {key: _column_index(header, key) for key in filters}
        values, included, missing = [], [], []
        units = _units(header)
        for number, row, cells in table['records']():
            if number <= table['header_row']:
                continue
            if _unit_row(row, number, table['header_row']):
                for h, raw in zip(header, row):
                    if h and isinstance(raw, str) and (m := UNIT.fullmatch(raw)):
                        units[h] = m.group(1)
                continue
            if any(not _accepts(_cell(row, cells, indices[key], number), condition) for key, condition in filters.items()):
                continue
            if any(isinstance(v, str) and re.fullmatch(r'\s*(?:总计|合计|小计|total|subtotal)\s*', v, re.I) for v in row):
                raise ValueError('Summary rows require explicit filters before aggregation')
            raw = _cell(row, cells, index, number)
            if raw in (None, ''):
                missing.append(number)
                continue
            if operation == 'count':
                value = 1
            else:
                if isinstance(raw, bool):
                    raise ValueError(f'Boolean is not a numeric measurement in row {number}')
                try:
                    value = float(raw)
                except (ValueError, TypeError):
                    raise ValueError(f'Non-numeric cell in row {number}') from None
                if not math.isfinite(value):
                    raise ValueError('Non-finite table value')
            values.append(value)
            included.append(number)
        if not values:
            raise ValueError('No matching numeric cells')
        functions = {'sum': lambda: math.fsum(values), 'mean': lambda: math.fsum(values) / len(values),
                     'min': lambda: min(values), 'max': lambda: max(values), 'count': lambda: len(values)}
        result = functions[operation]()
        if not math.isfinite(result):
            raise ValueError('Non-finite aggregate value')
        return {'operation': operation, 'column': column, 'sheet': table['sheet'], 'filters': filters,
                'value': result, 'rows': len(included), 'row_start': min(included), 'row_end': max(included),
                'included_rows': included, 'missing_rows': missing, 'header_row': table['header_row'],
                'unit': units.get(column), 'units': units,
                'note': 'Original units; blank target cells excluded; formula caches are required, never recalculated'}


def lookup(path, row, column=None, *, sheet=None):
    """Read one original physical row, never a header-relative data offset.

    Every returned cell has a coordinate, header and original unit. Blank cells
    are explicit. Uncached formulas are errors. Multiline CSV needs a structured
    selection so a logical record is never mislabeled as one physical line.
    """
    if isinstance(row, bool) or not isinstance(row, int) or not 1 <= row <= MAX_ROWS + 20:
        raise ValueError('Physical row must be a positive bounded integer')
    from openpyxl.utils import get_column_letter
    with _table(path, sheet) as table:
        header = table['header']
        indices = [_column_index(header, column)] if column is not None else list(range(len(header)))
        units = _units(header)
        selected = None
        row_values, unavailable = {}, {}
        for number, values, cells in table['records']():
            if _unit_row(values, number, table['header_row']):
                for h, raw in zip(header, values):
                    if h and isinstance(raw, str) and (m := UNIT.fullmatch(raw)):
                        units[h] = m.group(1)
            if number == row:
                selected = [(i, _cell(values, cells, i, number)) for i in indices]
                # Preserve row identity as well as the requested measurement.
                # Unrelated bad/formula cells cannot become invented values or
                # prevent a valid target lookup; their absence stays explicit.
                for i in range(len(header)):
                    coordinate = f'{get_column_letter(i + 1)}{row}'
                    try:
                        row_values[coordinate] = (i, _cell(values, cells, i, number))
                    except ValueError as exc:
                        unavailable[coordinate] = str(exc)
            if selected is not None and number >= max(row, table['header_row'] + 2):
                break
        if selected is None:
            raise ValueError('Physical row does not exist')
        output = {f'{get_column_letter(i + 1)}{row}': {'column': header[i],
                  'unit': units.get(header[i]), 'value': value} for i, value in selected}
        coordinates = list(output)
        return {'operation': 'lookup', 'column': column if column is not None else '*', 'sheet': table['sheet'],
                'filters': {}, 'value': next(iter(output.values()))['value'] if column is not None else [v['value'] for v in output.values()],
                'columns': [v['column'] for v in output.values() if v['column']],
                'values': [v['value'] for v in output.values()], 'cells': output,
                'rows': 1, 'row_start': row, 'row_end': row, 'included_rows': [row], 'missing_rows': [],
                'header_row': table['header_row'], 'unit': units.get(column), 'units': units,
                'range': f'{coordinates[0]}:{coordinates[-1]}',
                'row_context': {'range': f'A{row}:{get_column_letter(len(header))}{row}',
                    'sheet': table['sheet'], 'row': row,
                    'cells': {coordinate: {'column': header[i], 'unit': units.get(header[i]), 'value': value}
                              for coordinate, (i, value) in row_values.items()},
                    'unavailable_cells': unavailable},
                'note': 'One-based original physical row, not the Nth data record; units are not converted'}


CELL_REFERENCE = re.compile(r'(?<![A-Za-z0-9_])\$?([A-Za-z]{1,3})\$?([0-9]{1,8})(?![A-Za-z0-9_])')


def _simple_header_merges(path, worksheet_xml, header_row, width, column):
    """Inspect original XLSX merge metadata without loading a writable workbook.

    Check only the selected column's path. A full-width title above the header
    and a short, single-column vertical header merge have one explicit label.
    Horizontal groups or extra ancestors on that path require explicit structure.
    read_only worksheets expose their ZIP member path but not merged_cells.
    """
    from xml.etree.ElementTree import iterparse
    from zipfile import ZipFile
    from openpyxl.utils.cell import range_boundaries
    if not isinstance(worksheet_xml, str):
        raise ValueError('Worksheet merge metadata unavailable; cannot verify a simple header')
    header_merge = None
    with ZipFile(path) as archive, archive.open(worksheet_xml.lstrip('/')) as stream:
        for _, element in iterparse(stream, events=('end',)):
            if element.tag.rsplit('}', 1)[-1] == 'mergeCell':
                reference = element.attrib.get('ref', '')
                bounds = range_boundaries(reference)
                first_col, first_row, last_col, last_row = bounds
                if any(type(value) is not int for value in bounds):
                    raise ValueError('Invalid worksheet merge coordinates')
                if first_col <= column <= last_col and first_row <= header_row:
                    title = first_row == last_row < header_row and first_col == 1 and last_col >= width
                    vertical = (first_col == last_col == column and first_row == header_row
                                and header_row <= last_row <= header_row + 2)
                    if vertical:
                        if header_merge is not None:
                            raise ValueError('Overlapping merged table headers require explicit structure')
                        header_merge = reference
                    elif not title:
                        raise ValueError('Unsupported merged or multi-row table header in requested column; explicit structure is required')
            element.clear()
    return header_merge


def header_lookup(path, coordinate, *, sheet=None):
    """Use a physical anchor's column to read its detected original header cell.

    Duplicate labels remain unambiguous by column coordinate. The anchor is
    checked for existence/formula errors, but its data value is never the answer.
    Units come only from this column's header or an explicit nearby unit row.
    """
    from openpyxl.utils.cell import column_index_from_string, get_column_letter
    match = CELL_REFERENCE.fullmatch(coordinate) if isinstance(coordinate, str) else None
    if match is None:
        raise ValueError('Header lookup requires exactly one explicit cell coordinate')
    letter, number = match.groups()
    index, anchor_row = column_index_from_string(letter.upper()) - 1, int(number)
    if not 0 <= index < 16384 or not 1 <= anchor_row <= MAX_ROWS + 20:
        raise ValueError('Header coordinate is outside the supported physical range')
    with _table(path, sheet) as table:
        header, header_row = table['header'], table['header_row']
        if index >= len(header):
            raise ValueError('Header coordinate column does not exist')
        width = max((i + 1 for i, value in enumerate(header) if value), default=0)
        header_merge = None
        if Path(path).suffix.lower() == '.xlsx':
            header_merge = _simple_header_merges(path, table.get('worksheet_xml'), header_row, width, index + 1)
        anchor_found = False
        header_value = None
        unit = None
        unit_source = None
        for physical_row, values, cells in table['records']():
            if physical_row <= header_row:
                checked = [_cell(values, cells, i, physical_row) for i in range(len(header))]
                if physical_row == header_row:
                    filled = [value for value in checked if value not in (None, '')]
                    if not filled or any(not isinstance(value, str) for value in filled):
                        raise ValueError('Detected header is not a simple text row; explicit structure is required')
                    header_value = checked[index]
                    if not isinstance(header_value, str) or not header_value.strip():
                        raise ValueError('Requested column has no explicit header cell')
                    unit = _units([header_value]).get(header_value)
            if _unit_row(values, physical_row, header_row):
                raw = _cell(values, cells, index, physical_row)
                match_unit = UNIT.fullmatch(raw) if isinstance(raw, str) else None
                if match_unit:
                    current_unit = match_unit.group(1)
                    if unit is not None and current_unit != unit:
                        raise ValueError('Conflicting explicit units in requested header column')
                    unit = current_unit
                    unit_source = {'coordinate': f'{get_column_letter(index + 1)}{physical_row}',
                                   'row': physical_row, 'value': raw}
            if physical_row == anchor_row:
                _cell(values, cells, index, physical_row)
                anchor_found = True
            if anchor_found and physical_row >= max(anchor_row, header_row + 2):
                break
        if not anchor_found:
            raise ValueError('Header anchor physical row does not exist')
        if header_value is None:
            raise ValueError('Detected header row does not exist')
        letter = get_column_letter(index + 1)
        address = f'{letter}{header_row}'
        cell = {'column': header_value, 'unit': unit, 'value': header_value}
        if header_merge is not None:
            cell['merged_range'] = header_merge
        if unit_source is not None:
            cell['unit_source'] = unit_source
        return {'operation': 'header_lookup', 'column': header_value, 'sheet': table['sheet'],
                'requested_cell': f'{letter}{anchor_row}', 'column_letter': letter, 'column_index': index + 1,
                'value': header_value, 'values': [header_value], 'columns': [header_value],
                'cells': {address: cell}, 'range': f'{address}:{address}', 'filters': {},
                'rows': 1, 'row_start': header_row, 'row_end': header_row,
                'included_rows': [header_row], 'missing_rows': [], 'header_row': header_row,
                'header_merge_range': header_merge,
                'unit': unit, 'units': {header_value: unit} if unit is not None else {}, 'unit_source': unit_source,
                'note': 'Column selected by the explicit physical anchor; value is the detected original header cell, '
                        'not the anchor data. A short vertical header merge uses its original top cell. '
                        'Duplicate labels are resolved by column coordinate; original units are not converted.'}


def _annex_number(token):
    if token.isdigit():
        return int(token)
    digits = {'零': 0, '〇': 0, '一': 1, '二': 2, '三': 3, '四': 4, '五': 5,
              '六': 6, '七': 7, '八': 8, '九': 9}
    if token in digits:
        return digits[token]
    if '十' in token and token.count('十') == 1:
        tens, ones = token.split('十')
        if (not tens or tens in digits) and (not ones or ones in digits):
            return digits.get(tens, 1) * 10 + digits.get(ones, 0)
    return None


ANNEX = re.compile(r'附表\s*([0-9]+|[零〇一二三四五六七八九十]+)')


def _infer_sheet(question, names):
    aliases = {_annex_number(x) for x in ANNEX.findall(question)}
    if len(aliases) > 1 or None in aliases:
        return None
    exact = [name for name in names if re.search(re.escape(name) + (r'(?!\d)' if name[-1:].isdigit() else ''), question)]
    if len(exact) == 1:
        return exact[0]
    if len(exact) > 1:
        return None
    if aliases:
        if len(aliases) != 1 or None in aliases:
            return None
        matches = [name for name in names if any(_annex_number(x) in aliases for x in ANNEX.findall(name))]
        return matches[0] if len(matches) == 1 else None
    return names[0] if len(names) == 1 else None


def infer(question, path, *, diagnostics=None):
    """Infer unambiguous physical-row lookups or one-column bounded aggregates.

    Unsupported subsets return None; optional diagnostics records
    needs_structured_filter. Direct calculate() accepts explicit AND filters.
    """
    def decline(reason):
        if diagnostics is not None:
            diagnostics['reason'] = reason
        return None

    path = Path(path)
    # A workbook filename such as "附表1-12.xlsx" is not a sheet request.
    selection_question = question.replace(path.name, '')
    sheet = None
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = _infer_sheet(selection_question, book.sheetnames)
        finally:
            book.close()
        if sheet is None:
            return decline('ambiguous_sheet')
    elif path.suffix.lower() != '.csv':
        return decline('unsupported_format')
    with _table(path, sheet) as table:
        header = table['header']
    column_question = selection_question.replace(sheet, '') if sheet else selection_question
    column_question = ANNEX.sub(' ', column_question)
    mentioned = [h for h in header if h and re.search(
        (r'(?<![A-Za-z0-9_])' + re.escape(h) + r'(?![A-Za-z0-9_])') if h.isascii() else re.escape(h),
        column_question, re.I)]
    row_mentions = re.findall(r'第\s*(\d+)\s*行|\brow\s+(\d+)\b', question, re.I)
    operations = {'sum': ('总和', '合计', '求和', 'sum'), 'mean': ('平均', 'mean', 'average'),
                  'min': ('最小', 'minimum'), 'max': ('最大', 'maximum'), 'count': ('非空数量', 'count')}
    found = [op for op, words in operations.items() if any(
        re.search(r'\b' + w + r'\b', question, re.I) if w.isascii() else w in question for w in words)]
    residual = selection_question
    for words in operations.values():
        for word in words:
            residual = re.sub(r'\b' + word + r'\b', '', residual, flags=re.I) if word.isascii() else residual.replace(word, '')
    unsupported = r'其中|条件|大于|小于|超过|不足|至少|至多|等于|不等于|筛选|满足|排除|非零|非空|仅|只|且|同时|或者|之前|之后|以前|以后|以来|截至|截止|以外|不含|不包括|早于|晚于|最近|前\s*\d|后\s*\d|[<>=]|\b(?:between|where|and|or|if|filter|group|before|after|since|until|except)\b'
    if re.search(unsupported, residual, re.I):
        return decline('needs_structured_filter')
    header_question = bool(re.search(r'表头|列名|列标题|\b(?:column\s+(?:header|name|title)|header)\b', column_question, re.I))
    if header_question:
        coordinates = list(CELL_REFERENCE.finditer(column_question))
        if len(coordinates) != 1:
            return decline('needs_structured_filter')
        coordinate = coordinates[0]
        remainder = column_question[:coordinate.start()] + column_question[coordinate.end():]
        other_columns = re.findall(r'(?<![A-Za-z0-9_])([A-Za-z]{1,3})\s*列', remainder)
        column_range = re.search(r'(?<![A-Za-z0-9_])[A-Za-z]{1,3}\s*(?:[:：、,，/\-–—]|至|到|和|与|及)\s*[A-Za-z]{1,3}(?![A-Za-z0-9_])', remainder)
        if (row_mentions or found or column_range or any(c.upper() != coordinate[1].upper() for c in other_columns)
                or re.search(r'\d|分组|时段|月份|季度|站点|地区|按|为(?!何|什么|啥)|\b(?:during|when)\b', remainder, re.I)):
            return decline('needs_structured_filter')
        if diagnostics is not None:
            diagnostics['reason'] = 'physical_header_lookup'
        return header_lookup(path, coordinate[0], sheet=sheet)
    if row_mentions:
        if len(row_mentions) != 1 or found or re.search(r'数据(?:的)?第|第\s*\d+\s*条|数据行|至第|到第|前\s*\d|后\s*\d', question):
            return decline('needs_structured_filter')
        if re.search(r'\d{4}-\d{2}-\d{2}|\d{4}\s*年', question):
            return decline('needs_structured_filter')
        number = int(next(x for x in row_mentions[0] if x))
        if diagnostics is not None:
            diagnostics['reason'] = 'physical_row_lookup'
        return lookup(path, number, mentioned[0] if len(mentioned) == 1 else None, sheet=sheet)
    if len(mentioned) != 1 or len(found) != 1:
        return decline('ambiguous_column_or_operation')
    dates = re.findall(r'(?<!\d)\d{4}-\d{2}-\d{2}(?!\d)', question)
    years = re.findall(r'((?:19|20)\d{2})\s*年', question)
    date_columns = [h for h in header if h.casefold() in {'date', '日期'}]
    filters = {}
    if dates or years:
        if len(date_columns) != 1 or len(dates) > 2 or len(years) > 1 or dates and years:
            return decline('needs_structured_filter')
        try:
            for item in dates:
                date.fromisoformat(item)
        except ValueError:
            return decline('needs_structured_filter')
        if len(dates) == 1:
            filters = {date_columns[0]: dates[0]}
        elif len(dates) == 2:
            first = question.index(dates[0]) + len(dates[0])
            between = question[first:question.index(dates[1], first)]
            if not re.fullmatch(r'\s*(?:至|到|~|～|—|–|-|to|through)\s*', between, re.I) or dates[0] > dates[1]:
                return decline('needs_structured_filter')
            filters = {date_columns[0]: {'gte': dates[0], 'lte': dates[1]}}
        elif len(years) == 1:
            filters = {date_columns[0]: {'prefix': years[0] + '-'}}
    remaining = selection_question
    for item in dates + [y + '年' for y in years]:
        remaining = remaining.replace(item, '')
    if re.search(r'(?<!\d)(?:19|20)\d{2}(?!\d)|\d+\s*(?:月|日)|\d{1,2}[-/:]\d{1,2}|季度|月份|站点|地区|分组|时段|星期|周末|工作日', remaining):
        return decline('needs_structured_filter')
    if diagnostics is not None:
        diagnostics['reason'] = 'bounded_aggregate'
    return calculate(path, found[0], mentioned[0], filters=filters, sheet=sheet)
