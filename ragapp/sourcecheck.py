"""Validate physical citation locations against original files, independently of retrieval."""
import csv
import hashlib
import io
import json
from pathlib import Path
import re


def _metadata(path):
    """Read originals without invoking indexers, OCR or external models."""
    raw=path.read_bytes();ext=path.suffix.lower()
    data={'hash':hashlib.sha256(raw).hexdigest(),'extension':ext}
    if ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(''):raise ValueError('Encrypted PDF')
        data['pages']=len(reader.pages)
        # Most citations need only page bounds. Extract text lazily for the pages
        # whose new character spans can actually be checked, not whole textbooks.
        data['pdf_reader']=reader;data['page_text']={}
    elif ext=='.docx':
        from docx import Document
        doc=Document(path)
        data['paragraph_text']=[p.text for p in doc.paragraphs]
        data['paragraphs']=len(data['paragraph_text'])
        data['table_text']=[[' | '.join(c.text for c in row.cells) for row in table.rows] for table in doc.tables]
        data['tables']=[len(rows) for rows in data['table_text']]
        data['images']={i for i,r in enumerate(doc.part.rels.values()) if 'image' in r.reltype and not r.is_external}
    elif ext=='.xlsx':
        from openpyxl import load_workbook
        from .parsers import excel_header
        book=load_workbook(path,read_only=True,data_only=False)
        cached=load_workbook(path,read_only=True,data_only=True)
        data['sheets']={};data['sheet_rows']={};data['sheet_headers']={};data['sheet_columns']={}
        try:
            for sheet in book:
                data['sheets'][sheet.title]=(sheet.max_row,sheet.max_column)
                header_row,headers=excel_header(sheet)
                data['sheet_headers'][sheet.title]=headers
                data['sheet_columns'][sheet.title]=[str(c.value) for row in sheet.iter_rows(min_row=header_row,max_row=header_row) for c in row]
                units=[f'{c.coordinate}={c.value}' for row in sheet.iter_rows(min_row=header_row+1,max_row=min(header_row+2,sheet.max_row))
                       for c in row if isinstance(c.value,str) and re.fullmatch(r'\s*[（(][^）)]+[）)]\s*',c.value)]
                rows={}
                for number,(row,values) in enumerate(zip(sheet.iter_rows(),cached[sheet.title].iter_rows()),1):
                    fields=[]
                    for cell,value in zip(row,values):
                        if cell.value is None:continue
                        field=f'{cell.coordinate}={cell.value}'
                        if cell.data_type=='f':
                            field+=f' [formula cache={value.value!r}; '+('missing' if value.value is None else 'unverified freshness')+']'
                        fields.append(field)
                    rows[number]='Sheet: '+sheet.title+'\nHeaders: '+headers+'\nUnits: '+' | '.join(units)+'\n'+' | '.join(fields)
                data['sheet_rows'][sheet.title]=rows
        finally:book.close();cached.close()
    elif ext=='.xls':
        import xlrd
        book=xlrd.open_workbook(path)
        data['sheets']={s.name:(s.nrows,s.ncols) for s in book.sheets()}
        data['sheet_headers']={s.name:str(s.row_values(0)) for s in book.sheets()}
        data['sheet_rows']={s.name:{i+1:f'Headers: {s.row_values(0)}\n{s.row_values(i)}' for i in range(s.nrows)} for s in book.sheets()}
    elif ext in {'.png','.jpg','.jpeg','.webp','.bmp'}:
        from PIL import Image
        with Image.open(path) as image:image.verify()
        data['image']=True
    else:
        from .parsers import read_text
        text=read_text(path);data.update(text=text,lines=len(text.splitlines()))
        if ext=='.csv':
            rows=list(csv.reader(io.StringIO(text)));data['rows']=len(rows)
            data['columns']=rows[0] if rows else []
            data['csv_text']={i:' | '.join(f'{h}={v}' for h,v in zip(data['columns'],row)) for i,row in enumerate(rows[1:],2)}
    return data


def _range(start,end,maximum):
    return (isinstance(start,int) and not isinstance(start,bool) and isinstance(end,int)
            and not isinstance(end,bool) and 1<=start<=end<=maximum)


def _physical(location,data,path):
    errors=[]
    if 'page' in location:
        if not _range(location['page'],location['page'],data.get('pages',0)):errors.append('invalid_page')
        if 'bbox' in location:
            box=location['bbox']
            if (not isinstance(box,(list,tuple)) or len(box)!=4 or
                not all(isinstance(n,(int,float)) and not isinstance(n,bool) for n in box) or
                not 0<=box[0]<box[2]<=1 or not 0<=box[1]<box[3]<=1):errors.append('invalid_bbox')
    elif 'paragraph' in location or 'paragraph_start' in location:
        first=location.get('paragraph_start',location.get('paragraph'))
        last=location.get('paragraph_end',location.get('paragraph',first))
        if not _range(first,last,data.get('paragraphs',0)):errors.append('invalid_paragraph')
    elif 'table' in location:
        table=location['table'];tables=data.get('tables',[])
        if not _range(table,table,len(tables)):errors.append('invalid_table_row')
        else:
            first=location.get('row_start',location.get('row',0));last=location.get('row_end',location.get('row',first))
            if not _range(first,last,tables[table-1]):errors.append('invalid_table_row')
    elif 'embedded_image' in location:
        if location['embedded_image'] not in data.get('images',set()):errors.append('invalid_embedded_image')
    elif location.get('sheet') is not None or data['extension'] in {'.xlsx','.xls'}:
        dimensions=data.get('sheets',{}).get(location.get('sheet'))
        if not dimensions:errors.append('invalid_sheet')
        else:
            positioned=False
            if location.get('range'):
                from openpyxl.utils.cell import range_boundaries
                positioned=True
                try:
                    c1,r1,c2,r2=range_boundaries(location['range'])
                    if not (_range(r1,r2,dimensions[0]) and _range(c1,c2,dimensions[1])):errors.append('invalid_cell_range')
                    elif 'row' in location and not r1<=location['row']<=r2:errors.append('inconsistent_sheet_row_locator')
                except (TypeError,ValueError):errors.append('invalid_cell_range')
            if 'row_start' in location or 'row_end' in location:
                positioned=True
                if not _range(location.get('row_start'),location.get('row_end'),dimensions[0]):errors.append('invalid_sheet_row_range')
            elif 'row' in location:
                positioned=True
                if not _range(location['row'],location['row'],dimensions[0]):errors.append('invalid_sheet_row')
            if location.get('kind')=='calculation':
                names=data.get('sheet_columns',{}).get(location['sheet'],[])
                lookup=location.get('operation')=='lookup'
                if lookup and location.get('column')=='*':
                    if not location.get('columns') or any(c not in names for c in location['columns']):errors.append('invalid_calculation_column')
                elif location.get('column') not in names:errors.append('invalid_calculation_column')
                if location.get('operation') not in {'sum','mean','min','max','count','lookup'}:errors.append('invalid_calculation_operation')
                if lookup:
                    if location.get('row_start')!=location.get('row_end'):errors.append('lookup_requires_one_physical_row')
                    if not location.get('cells'):errors.append('missing_lookup_cells')
                    for cell in location.get('cells',[]):
                        from openpyxl.utils.cell import coordinate_to_tuple
                        try:
                            row,column=coordinate_to_tuple(cell)
                            if row!=location.get('row_start') or not _range(column,column,dimensions[1]):errors.append('invalid_lookup_cell')
                        except (ValueError,TypeError,KeyError):errors.append('invalid_lookup_cell')
            if not positioned:errors.append('missing_physical_locator')
    elif 'row' in location or 'row_start' in location:
        first=location.get('row_start',location.get('row'));last=location.get('row_end',location.get('row',first))
        if not _range(first,last,data.get('rows',0)):errors.append('invalid_csv_range' if 'row_start' in location else 'invalid_csv_row')
    elif 'line_start' in location:
        if not _range(location['line_start'],location.get('line_end'),data.get('lines',0)):errors.append('invalid_line_range')
        if any(not _range(n,n,data.get('lines',0)) for n in location.get('matched_lines',{}).values()):errors.append('invalid_matched_line')
    elif 'image' in location:
        if location['image']!=path.name or not data.get('image'):errors.append('invalid_image')
    else:errors.append('missing_physical_locator')
    return errors


def _source_text(span,data):
    """Return exact parser serialization where it is verifiable without OCR."""
    if 'document_char_start' in span and 'text' in data:return data['text'],'document_char_start','document_char_end'
    if 'paragraph' in span and 'paragraph_text' in data:
        return data['paragraph_text'][span['paragraph']-1],'char_start','char_end'
    if 'table' in span and 'table_text' in data:
        return data['table_text'][span['table']-1][span['row']-1],'char_start','char_end'
    if span.get('sheet') in data.get('sheets',{}):
        text=(data['sheet_headers'][span['sheet']] if span.get('kind')=='table_schema' else
              data['sheet_rows'][span['sheet']].get(span.get('row')))
        return text,'char_start','char_end'
    if data['extension']=='.csv' and 'row' in span:
        text=' | '.join(data['columns']) if span.get('kind')=='table_schema' else data['csv_text'].get(span['row'])
        return text,'char_start','char_end'
    if span.get('method')=='pdf_text' and 'page' in span:
        page=span['page']
        if page not in data['page_text']:
            data['page_text'][page]=data['pdf_reader'].pages[page-1].extract_text() or ''
        return data['page_text'][page],'char_start','char_end'
    return None,'char_start','char_end'


def _verify_spans(location,citation,data,path):
    spans=location.get('source_spans')
    if location.get('context_expanded') and location.get('parent_source_spans'):
        spans=location['parent_source_spans']
    elif spans is None and location.get('parent_source_spans'):
        spans=location['parent_source_spans']
    if spans is None:return None
    if not isinstance(spans,list) or not spans:return [{'valid':False,'errors':['invalid_source_spans'],'text_verified':False}]
    text=citation.get('text');checked=[];previous_end=0
    for number,span in enumerate(spans):
        errors=[];verified=None
        if not isinstance(span,dict):
            checked.append({'valid':False,'errors':['invalid_source_span'],'text_verified':False});continue
        if span.get('path',citation['path']).replace('\\','/')!=citation['path'].replace('\\','/'):
            errors.append('span_source_mismatch')
        if span.get('version',citation['version'])!=citation['version']:errors.append('span_version_mismatch')
        errors.extend(_physical(span,data,path))
        a,b=span.get('chunk_start'),span.get('chunk_end')
        valid_offsets=isinstance(a,int) and not isinstance(a,bool) and isinstance(b,int) and not isinstance(b,bool) and 0<=a<b
        if not valid_offsets or (isinstance(text,str) and b>len(text)):
            errors.append('invalid_chunk_character_range')
        elif isinstance(text,str):
            if a<previous_end:errors.append('overlapping_source_spans')
            elif text[previous_end:a].strip():errors.append('uncited_text_between_spans')
            previous_end=b
            try:
                source,start_key,end_key=_source_text(span,data) if not errors else (None,'char_start','char_end')
                first,last=span.get(start_key),span.get(end_key)
                if not (isinstance(first,int) and not isinstance(first,bool) and isinstance(last,int) and not isinstance(last,bool) and 0<=first<last):
                    errors.append('invalid_source_character_range')
                elif last-first!=b-a:errors.append('source_span_length_mismatch')
                elif source is not None:
                    if last>len(source):errors.append('invalid_source_character_range')
                    elif source[first:last]!=text[a:b]:errors.append('source_text_mismatch')
                    else:
                        verified=True
                        if start_key=='document_char_start':
                            if span.get('line_start')!=source[:first].count('\n')+1 or span.get('line_end')!=source[:last-1].count('\n')+1:
                                errors.append('source_line_character_mismatch')
            except (IndexError,KeyError,TypeError):errors.append('invalid_source_character_range')
        checked.append({'index':number,'locator':span,'valid':not errors,'errors':errors,
                        'text_verified':False if errors else verified})
    if isinstance(text,str) and text[previous_end:].strip():
        checked.append({'valid':False,'errors':['uncited_text_after_spans'],'text_verified':False})
    return checked


def validate(report_path,corpus,snapshot):
    """Check frozen versions and every cited span; retain legacy locations.

    Text equality is separate from physical validity. Layout/image OCR whose
    content cannot be independently reconstructed is marked for content review.
    """
    report=json.loads(Path(report_path).read_text(encoding='utf-8'))
    corpus=Path(corpus).resolve();cache={};results=[]
    normalized_snapshot={name.replace('\\','/'):value for name,value in snapshot.items()}
    for row in report['results']:
        checked=[]
        for citation in row['answer'].get('citations',[]):
            name=citation['path'];location=citation['locator'];path=(corpus/name).resolve()
            if isinstance(location,str):location=json.loads(location)
            errors=[];spans=None
            if not path.is_relative_to(corpus) or not path.is_file():errors.append('source_missing_or_outside_root')
            else:
                if name not in cache:cache[name]=_metadata(path)
                data=cache[name];expected=normalized_snapshot.get(name.replace('\\','/'))
                if expected is None or expected['hash']!=data['hash'] or expected['version']!=citation['version']:
                    errors.append('source_version_mismatch')
                spans=_verify_spans(location,citation,data,path)
                if spans is None:errors.extend(_physical(location,data,path))
                else:
                    for span in spans:errors.extend(span['errors'])
            checked.append({'citation':citation['citation'],'path':name,'locator':location,'valid':not errors,
                'errors':list(dict.fromkeys(errors)),'spans':spans,
                'text_verified':(False if any(s['text_verified'] is False for s in spans) else
                                 None if any(s['text_verified'] is None for s in spans) else True) if spans else None,
                'requires_content_review':bool(spans and any(s['text_verified'] is None for s in spans))})
        results.append({'id':row['id'],'valid':all(c['valid'] for c in checked),'citations':checked})
    return {'all_valid':all(r['valid'] for r in results),'citation_count':sum(len(r['citations']) for r in results),'results':results}

