import copy
import hashlib
import json

import pytest

from ragapp.parsers import parse, chunk
from ragapp.context import build_context, public_evidence
from ragapp.sourcecheck import validate


class Tokenizer:
    def __call__(self,text,**kwargs):
        return {'offset_mapping':[(i,i+1) for i in range(len(text))]}


def check(tmp_path,path,citations):
    report=tmp_path/'report.json'
    report.write_text(json.dumps({'results':[{'id':'q1','answer':{'citations':citations}}]}),encoding='utf-8')
    snapshot={path.name:{'hash':hashlib.sha256(path.read_bytes()).hexdigest(),'version':'v'}}
    return validate(report,path.parent,snapshot)


def citation(path,location,text=None):
    result={'citation':'S1','path':path.name,'version':'v','locator':location}
    if text is not None:result['text']=text
    return result


def test_second_span_must_be_checked_even_when_top_paragraph_is_valid(tmp_path):
    from docx import Document
    path=tmp_path/'source.docx';doc=Document();doc.add_paragraph('Alpha');doc.add_paragraph('Beta');doc.save(path)
    item=list(chunk(parse(path,tmp_path/'cache'),Tokenizer(),100,8))[0]
    good=citation(path,item['locator'],item['text'])
    assert check(tmp_path,path,[good])['all_valid']
    bad=copy.deepcopy(good);bad['locator']['source_spans'][1]['paragraph']=20
    result=check(tmp_path,path,[bad])
    assert not result['all_valid']
    assert 'invalid_paragraph' in result['results'][0]['citations'][0]['errors']


def test_character_and_line_offsets_reject_wrong_text_and_wrong_line(tmp_path):
    path=tmp_path/'source.py';path.write_bytes(b'def function():\r\n    value = 123\r\n    return value\r\n')
    items=list(chunk(parse(path,tmp_path/'cache'),Tokenizer(),25,5))
    citations=[citation(path,item['locator'],item['text']) for item in items]
    result=check(tmp_path,path,citations)
    assert result['all_valid']
    assert all(c['text_verified'] for c in result['results'][0]['citations'])
    wrong=copy.deepcopy(citations[0]);wrong['text']=wrong['text'].replace('def','bad')
    assert not check(tmp_path,path,[wrong])['all_valid']
    wrong=copy.deepcopy(citations[-1]);wrong['locator']['source_spans'][0]['line_start']=1
    result=check(tmp_path,path,[wrong])
    assert not result['all_valid']
    assert 'source_line_character_mismatch' in result['results'][0]['citations'][0]['errors']


def test_expanded_parent_spans_are_verified_against_expanded_text(tmp_path):
    path=tmp_path/'source.txt';path.write_bytes(b'First sentence. Second sentence. Third sentence. Fourth sentence.')
    items=list(chunk(parse(path,tmp_path/'cache'),Tokenizer(),30,4))
    hit={**items[1],'id':'c','document_id':'d','version':'v','path':path.name,'citation':'S1'}
    selected,_=build_context([hit],Tokenizer(),10000,True)
    assert selected[0]['locator']['context_expanded']
    assert check(tmp_path,path,[public_evidence(selected[0])])['all_valid']
    # Internal representations may provide parent_source_spans directly.
    internal=copy.deepcopy(selected[0]);internal['locator'].pop('source_spans')
    assert check(tmp_path,path,[internal])['all_valid']


def test_uncited_content_and_foreign_spans_cannot_hide_inside_citation(tmp_path):
    path=tmp_path/'source.txt';path.write_text('Original evidence.')
    item=list(chunk(parse(path,tmp_path/'cache'),Tokenizer(),100,8))[0]
    good=citation(path,item['locator'],item['text'])
    bad=copy.deepcopy(good);bad['text']+=' fabricated extra fact'
    result=check(tmp_path,path,[bad])
    assert 'uncited_text_after_spans' in result['results'][0]['citations'][0]['errors']
    bad=copy.deepcopy(good);bad['locator']['source_spans'][0]['path']='other.txt'
    assert not check(tmp_path,path,[bad])['all_valid']


def test_xlsx_schema_rows_and_calculation_ranges(tmp_path):
    from openpyxl import Workbook
    path=tmp_path/'source.xlsx';book=Workbook();sheet=book.active;sheet.title='Inventory'
    sheet.append(['device','price']);sheet.append(['A',10]);sheet.append(['B',20]);book.save(path)
    items=list(chunk(parse(path,tmp_path/'cache'),Tokenizer(),150,8))
    assert check(tmp_path,path,[citation(path,c['locator'],c['text']) for c in items])['all_valid']
    loc={'sheet':'Inventory','row_start':2,'row_end':3,'column':'price','operation':'sum','kind':'calculation'}
    assert check(tmp_path,path,[citation(path,loc)])['all_valid']
    result=check(tmp_path,path,[citation(path,{**loc,'row_end':300})])
    assert 'invalid_sheet_row_range' in result['results'][0]['citations'][0]['errors']
    assert not check(tmp_path,path,[citation(path,{**loc,'sheet':'Missing'})])['all_valid']
    assert not check(tmp_path,path,[citation(path,{**loc,'column':'made_up'})])['all_valid']


def test_legacy_and_csv_calculator_sheet_none_are_compatible(tmp_path):
    path=tmp_path/'source.csv';path.write_text('day,value\na,1\nb,2\n')
    assert check(tmp_path,path,[citation(path,{'row':2})])['all_valid']
    loc={'sheet':None,'row_start':2,'row_end':3,'column':'value','operation':'sum','kind':'calculation'}
    assert check(tmp_path,path,[citation(path,loc)])['all_valid']
    assert not check(tmp_path,path,[citation(path,{**loc,'row_end':4})])['all_valid']


def test_layout_ocr_valid_coordinates_do_not_claim_text_verified(tmp_path):
    from pypdf import PdfWriter
    path=tmp_path/'source.pdf';writer=PdfWriter();writer.add_blank_page(width=600,height=800)
    with path.open('wb') as stream:writer.write(stream)
    span={'page':1,'bbox':[0,.1,.5,.2],'method':'mineru','block':1,
          'char_start':0,'char_end':8,'chunk_start':0,'chunk_end':8}
    result=check(tmp_path,path,[citation(path,{'source_spans':[span]},'OCR text')])
    checked=result['results'][0]['citations'][0]
    assert checked['valid'] and not checked['text_verified'] and checked['requires_content_review']
    result=check(tmp_path,path,[citation(path,{'source_spans':[{**span,'page':2}]},'OCR text')])
    assert not result['all_valid']


def literal_pair(text,spans,quote=None,locator=None):
    return ({'path':'source','locator':locator or {'page':2},'span_text':quote or text},
            {'path':'source','version':'v','text':text,'locator':{'source_spans':spans}})


def test_literal_gold_can_cross_contiguous_char_spans_on_one_page():
    from ragapp.evaluation import span_matches
    gold,item=literal_pair('Voltage is 12 mV.',[
        {'page':2,'method':'pdf_text','chunk_start':0,'chunk_end':11,'char_start':40,'char_end':51},
        {'page':2,'method':'pdf_text','chunk_start':11,'chunk_end':17,'char_start':51,'char_end':57}])
    assert span_matches(gold,item)
    item['locator']['source_spans'][1]['char_start']=52
    assert not span_matches(gold,item)  # A known source-character gap is not a whitespace join.


def test_literal_gold_crosses_adjacent_same_section_paragraphs_and_lines():
    from ragapp.evaluation import span_matches
    gold,item=literal_pair('First fact.\n\nSecond fact.',[
        {'paragraph':3,'section_path':['Results'],'chunk_start':0,'chunk_end':11,'char_start':0,'char_end':11},
        {'paragraph':4,'section_path':['Results'],'chunk_start':13,'chunk_end':25,'char_start':0,'char_end':12}],
        locator={'paragraph_start':3,'paragraph_end':4,'section_path':['Results']})
    assert span_matches(gold,item)
    item['locator']['source_spans'][1]['paragraph']=5
    assert not span_matches(gold,item)
    gold,item=literal_pair('alpha\nbeta',[
        {'line_start':10,'line_end':10,'chunk_start':0,'chunk_end':6,'document_char_start':90,'document_char_end':96},
        {'line_start':11,'line_end':11,'chunk_start':6,'chunk_end':10,'document_char_start':96,'document_char_end':100}],
        locator={'line_start':10,'line_end':11})
    assert span_matches(gold,item)
    gold['locator']={'line':10}
    assert not span_matches(gold,item)


def test_same_page_adjacent_parser_blocks_can_share_a_literal_quote():
    from ragapp.evaluation import span_matches
    gold,item=literal_pair('Heading\n\nBody text',[
        {'page':2,'block':7,'section_path':['Section'],'chunk_start':0,'chunk_end':7,'char_start':0,'char_end':7},
        {'page':2,'block':8,'section_path':['Section'],'chunk_start':9,'chunk_end':18,'char_start':0,'char_end':9}])
    assert span_matches(gold,item)
    item['locator']['source_spans'][1]['section_path']=['Different section']
    assert not span_matches(gold,item)


def test_cross_page_quote_cannot_borrow_matching_top_locator_or_expected_page():
    from ragapp.evaluation import span_matches
    gold,item=literal_pair('alpha beta',[
        {'page':2,'chunk_start':0,'chunk_end':5},
        {'page':3,'chunk_start':6,'chunk_end':10}])
    item['locator']['page']=2
    assert not span_matches(gold,item)
    gold['locator']={'page':3}
    assert not span_matches(gold,item)


def test_each_repeated_literal_occurrence_uses_its_own_physical_char_interval():
    from ragapp.evaluation import span_matches
    gold,item=literal_pair('alpha beta alpha beta',[
        {'page':1,'chunk_start':0,'chunk_end':10},
        {'page':2,'chunk_start':11,'chunk_end':16},
        {'page':3,'chunk_start':17,'chunk_end':21}],quote='alpha beta')
    assert not span_matches(gold,item)
    gold['locator']={'page':1}
    assert span_matches(gold,item)
    # A later complete occurrence can match even when the first one is on another page.
    item['locator']['source_spans']=[{'page':1,'chunk_start':0,'chunk_end':10},
                                   {'page':2,'chunk_start':11,'chunk_end':21}]
    gold['locator']={'page':2}
    assert span_matches(gold,item)


def test_unmapped_body_parent_and_unknown_same_page_joins_are_not_literal_hits():
    from ragapp.evaluation import span_matches
    gold,item=literal_pair('alpha beta',[{'page':2,'chunk_start':0,'chunk_end':5},
                                       {'page':2,'chunk_start':6,'chunk_end':10}])
    assert not span_matches(gold,item)  # Same page does not establish adjacency.
    item['locator']={'page':2,'context_expanded':True,'parent_source_spans':[{'page':2}]}
    assert not span_matches(gold,item)
    item['locator']={'page':2,'source_spans':[]}
    assert not span_matches(gold,item)
    item['locator']={'page':2};item['text']='unrelated';item['parent']='alpha beta'
    assert not span_matches(gold,item)
    item['text']='alpha beta';item['locator']={'source_spans':[{'symbol':'f','chunk_start':0,'chunk_end':10}]}
    gold['locator']={'symbol':'f'}
    assert not span_matches(gold,item)  # A function name alone is not a physical mapping.


def test_nonwhitespace_visible_gaps_cannot_be_ignored_even_when_gold_matches_full_text():
    from ragapp.evaluation import span_matches
    gold,item=literal_pair('alpha NOT beta',[
        {'page':2,'block':1,'chunk_start':0,'chunk_end':5},
        {'page':2,'block':2,'chunk_start':10,'chunk_end':14}])
    assert not span_matches(gold,item)
    gold['span_text']='alpha beta'
    assert not span_matches(gold,item)


@pytest.mark.parametrize('quote',['Power is 42 MW.','Power is 42 mW!','Power is not 42 mW.','x2 = 4'])
def test_literal_units_symbols_punctuation_and_negation_are_preserved(quote):
    from ragapp.evaluation import span_matches
    text='Power is 42 mW. x² = 4'
    gold,item=literal_pair(text,[{'page':2,'chunk_start':0,'chunk_end':len(text)}],quote=quote)
    assert not span_matches(gold,item)


def test_nfc_mapping_keeps_correct_offsets_without_formula_or_case_normalization():
    from ragapp.evaluation import span_matches
    text='Wrong. Cafe\u0301 is 2 mW.'
    gold,item=literal_pair(text,[{'page':1,'chunk_start':0,'chunk_end':7},
                                {'page':2,'chunk_start':7,'chunk_end':len(text)}],quote='Café\n is 2 mW.')
    assert span_matches(gold,item)
    gold['locator']={'page':1}
    assert not span_matches(gold,item)


@pytest.mark.parametrize('mutation',[
    lambda spans:spans[0].update(chunk_start=True),
    lambda spans:spans[0].pop('chunk_end'),
    lambda spans:spans[0].update(path='other'),
    lambda spans:spans[0].update(version='other'),
    lambda spans:spans.append(dict(spans[0])),
])
def test_literal_mapping_rejects_ambiguous_offsets_foreign_sources_and_overlap(mutation):
    from ragapp.evaluation import span_matches
    gold,item=literal_pair('alpha',[{'page':2,'chunk_start':0,'chunk_end':5}])
    mutation(item['locator']['source_spans'])
    assert not span_matches(gold,item)
