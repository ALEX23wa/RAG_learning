"""Original table fixtures: no embedding models or real index writes."""
from datetime import datetime
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from ragapp.calculator import calculate, header_lookup, infer, lookup, prompt_calculation


@pytest.fixture
def inventory(tmp_path):
    path = tmp_path / '附表1-12.xlsx'
    book = Workbook()
    ws = book.active
    ws.title = '附表3-设备原值统计'
    ws.append(['设备台账'])
    ws.append(['设备', '原值', '日期'])
    ws.append([None, '（万元）', None])
    ws.append(['传感器', 2, datetime(2024, 1, 1)])
    ws.append(['电话', 3, datetime(2024, 1, 2)])
    other = book.create_sheet('附表13-其他')
    other.append(['设备', '原值'])
    other.append(['错误表设备', 900])
    book.save(path)
    book.close()
    return path


@pytest.fixture
def daily(tmp_path):
    path = tmp_path / 'daily.csv'
    path.write_text('date,kind,e,tp\n2024-01-01,a,1,2\n2024-01-02,b,2,4\n2024-01-03,a,3,8\n2025-01-01,a,4,16\n', encoding='utf-8')
    return path


def test_lookup_uses_original_physical_row_and_units(inventory):
    result = infer('附表3第 4 行的原值是多少？', inventory)
    assert result['value'] == 2
    assert result['operation'] == 'lookup'
    assert result['row_start'] == result['row_end'] == 4
    assert result['sheet'] == '附表3-设备原值统计'
    assert result['header_row'] == 2
    assert result['unit'] == '万元'
    assert result['range'] == 'B4:B4'
    assert result['cells'] == {'B4': {'column': '原值', 'unit': '万元', 'value': 2}}
    assert result['row_context']['cells']['A4']['value'] == '传感器'
    assert result['row_context']['cells']['C4']['value'] == '2024-01-01T00:00:00'
    assert result['row_context']['range'] == 'A4:C4'
    assert result['row_context']['sheet'] == result['sheet']


def test_lookup_does_not_invent_unrelated_uncached_formula_values(inventory):
    from openpyxl import load_workbook
    book = load_workbook(inventory)
    book.active['C4'] = '=1+1'
    book.save(inventory)
    book.close()
    result = infer('附表3第4行原值', inventory)
    assert result['value'] == 2
    assert result['row_context']['cells']['A4']['value'] == '传感器'
    assert 'C4' not in result['row_context']['cells']
    assert 'Missing formula cache' in result['row_context']['unavailable_cells']['C4']
    with pytest.raises(ValueError, match='Missing formula cache'):
        lookup(inventory, 4, '日期', sheet='附表3-设备原值统计')


def test_aggregate_prompt_compresses_large_membership_without_losing_holes(tmp_path):
    import json
    path = tmp_path/'many.csv'
    path.write_text('value\n' + ''.join(('\n' if i % 100 == 0 else '2\n') for i in range(9000)))
    result = calculate(path, 'mean', 'value')
    original = json.dumps(result)
    compact = prompt_calculation(result)
    assert compact['value'] == 2 and compact['rows'] == 8910
    for key in ('included_rows', 'missing_rows'):
        restored = [row for start, end in compact[key + '_ranges'] for row in range(start, end+1)]
        assert restored == result[key]
        assert compact[key + '_count'] == len(result[key])
    assert len(json.dumps(compact)) < 5000
    assert json.dumps(result) == original


def test_lookup_whole_row_keeps_all_coordinates(inventory):
    result = infer('请给出附表三第4行的内容', inventory)
    assert result['column'] == '*'
    assert result['columns'] == ['设备', '原值', '日期']
    assert result['values'] == ['传感器', 2, '2024-01-01T00:00:00']
    assert result['range'] == 'A4:C4'
    assert set(result['cells']) == {'A4', 'B4', 'C4'}
    assert infer('附表3第4行的设备和原值', inventory)['column'] == '*'


def test_lookup_headers_units_and_missing_rows(inventory):
    sheet = '附表3-设备原值统计'
    assert lookup(inventory, 2, '原值', sheet=sheet)['value'] == '原值'
    assert lookup(inventory, 3, '原值', sheet=sheet)['value'] == '（万元）'
    with pytest.raises(ValueError, match='does not exist'):
        lookup(inventory, 40, sheet=sheet)
    with pytest.raises(ValueError, match='positive bounded'):
        lookup(inventory, 0, sheet=sheet)


def test_csv_lookup_is_not_data_relative(daily):
    assert infer('第2行 tp 列是多少', daily)['value'] == '2'
    assert infer('第1行内容', daily)['values'] == ['date', 'kind', 'e', 'tp']
    assert infer('数据的第2行 tp 是多少', daily) is None
    assert infer('第2行和第3行 tp 是多少', daily) is None
    assert infer('2024-01-01第2行 tp', daily) is None


def test_annex_disambiguation_and_workbook_name(inventory):
    assert infer('附表1-12.xlsx 中附表3的原值合计', inventory)['value'] == 5
    assert infer('附表十三的原值合计', inventory)['value'] == 900
    assert infer('原值合计', inventory) is None
    assert infer('附表3和附表13原值合计', inventory) is None
    assert infer('附表30原值合计', inventory) is None
    from openpyxl import load_workbook
    book = load_workbook(inventory)
    book.create_sheet('附表3-另一版本').append(['设备', '原值'])
    book.save(inventory)
    book.close()
    assert infer('附表3原值合计', inventory) is None
    assert infer('附表3-设备原值统计 原值合计', inventory)['value'] == 5


def test_numeric_sheet_prefix_is_not_exact_match(tmp_path):
    path = tmp_path / 'sheets.xlsx'
    wb = Workbook()
    wb.active.title = '附表3'
    wb.active.append(['name', 'value'])
    wb.active.append(['wrong', 2])
    ws = wb.create_sheet('附表30')
    ws.append(['name', 'value'])
    ws.append(['correct', 30])
    wb.save(path)
    wb.close()
    assert infer('附表30 value sum', path)['value'] == 30


def test_explicit_and_filters_and_exact_included_rows(daily):
    result = calculate(daily, 'sum', 'tp', filters={'kind': 'a', 'date': {'gte': '2024-01-01', 'lte': '2024-12-31'}})
    assert result['value'] == 10
    assert result['included_rows'] == [2, 4]
    assert result['row_start'] == 2 and result['row_end'] == 4
    assert calculate(daily, 'mean', 'tp', filters={'kind': 'a'})['value'] == pytest.approx(26 / 3)
    assert calculate(daily, 'count', 'kind')['value'] == 4


def test_supported_date_queries_and_excel_date_cells(daily, inventory):
    assert infer('全部 tp 列的总和', daily)['value'] == 30
    assert infer('2024年 tp 平均值', daily)['value'] == pytest.approx(14 / 3)
    assert infer('2024-01-02 tp 总和', daily)['value'] == 4
    assert infer('2024-01-01至2024-01-02 tp 总和', daily)['value'] == 6
    assert infer('2024-01-01 to 2024-01-02 tp SUM', daily)['value'] == 6
    assert infer('附表3 2024-01-01 原值合计', inventory)['value'] == 2
    assert infer('附表3 2024年 原值合计', inventory)['value'] == 5


@pytest.mark.parametrize('question', [
    '2024-01-01和2024-01-03 tp 总和',
    '2024-01-01、2024-01-02、2024-01-03 tp 总和',
    '2024-01-03至2024-01-01 tp 总和',
    '2024年和2025年 tp 总和',
    '2024年 2024-01-01 tp 总和',
    '2024-02-31 tp 总和',
    '2024-01-01 tp 大于1的总和',
    '2024-01-01 tp 总和 where kind=a',
    '2024年 tp 总和且只保留非零值',
    '2024年1月 tp 总和',
    '2024-01-01之后 tp 总和',
    '2024年以外 tp 总和',
    '2024-01-01至01-03 tp 总和',
    '2024-01-01 12:00 tp 总和',
    '2024年第一季度 tp 总和',
    'tp 前3行总和',
    '第2行 tp 大于1时的值',
])
def test_unrepresentable_filters_are_not_silently_dropped(daily, question):
    diagnostics = {}
    assert infer(question, daily, diagnostics=diagnostics) is None
    assert diagnostics['reason'] == 'needs_structured_filter'


def test_date_requests_without_a_date_column_decline(tmp_path):
    path = tmp_path / 'a.csv'
    path.write_text('name,value\na,2\n', encoding='utf-8')
    assert infer('2024年 value 总和', path) is None


def test_formula_cache_errors_and_filtered_out_targets(inventory):
    from openpyxl import load_workbook
    book = load_workbook(inventory)
    ws = book['附表3-设备原值统计']
    ws['B5'] = '=1+2'
    book.save(inventory)
    book.close()
    with pytest.raises(ValueError, match='formula cache at B5'):
        infer('附表3第5行原值', inventory)
    with pytest.raises(ValueError, match='formula cache at B5'):
        calculate(inventory, 'sum', '原值', sheet=ws.title)
    assert calculate(inventory, 'sum', '原值', sheet=ws.title, filters={'设备': '传感器'})['value'] == 2
    book = load_workbook(inventory)
    book[ws.title]['A4'] = '=1'
    book.save(inventory)
    book.close()
    with pytest.raises(ValueError, match='formula cache at A4'):
        calculate(inventory, 'sum', '原值', sheet=ws.title, filters={'设备': '传感器'})


def test_real_xlsx_cached_formula_fixture(inventory):
    # openpyxl deliberately does not calculate formulas. Supply a fake cached
    # value in the real ZIP/XML format and verify that it is read, never eval'd.
    from openpyxl import load_workbook
    book = load_workbook(inventory)
    book['附表3-设备原值统计']['B5'] = '=1+2'
    book.save(inventory)
    book.close()
    with ZipFile(inventory) as src:
        members = {info.filename: src.read(info.filename) for info in src.infolist()}
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    root = ElementTree.fromstring(members['xl/worksheets/sheet1.xml'])
    cell = root.find('.//s:c[@r="B5"]', ns)
    cell.find('s:v', ns).text = '3'
    members['xl/worksheets/sheet1.xml'] = ElementTree.tostring(root)
    with ZipFile(inventory, 'w') as dst:
        for name, content in members.items():
            dst.writestr(name, content)
    assert infer('附表3第5行原值', inventory)['value'] == 3
    assert infer('附表3原值合计', inventory)['value'] == 5


def test_missing_and_short_rows_are_not_index_errors(tmp_path):
    path = tmp_path / 'short.csv'
    path.write_text('name,value,kind\na,2,x\nb,3\nc,,x\n', encoding='utf-8')
    result = calculate(path, 'sum', 'value', filters={'kind': 'x'})
    assert result['value'] == 2
    assert result['missing_rows'] == [4]


def test_duplicate_headers_and_bad_filters_are_rejected(tmp_path):
    path = tmp_path / 'bad.csv'
    path.write_text('value,value\n1,2\n', encoding='utf-8')
    with pytest.raises(ValueError, match='unique'):
        calculate(path, 'sum', 'value')
    with pytest.raises(ValueError, match='unique'):
        lookup(path, 2, 'value')
    path.write_text('name,value\n', encoding='utf-8')
    with pytest.raises(ValueError, match='Unsupported filter operator'):
        calculate(path, 'sum', 'value', filters={'name': {'exec': 'bad'}})


@pytest.mark.parametrize('value', ['NaN', 'inf', '-inf', 'not-a-number'])
def test_non_numeric_aggregate_is_not_fabricated(tmp_path, value):
    path = tmp_path / 'bad.csv'
    path.write_text(f'name,value\na,{value}\n', encoding='utf-8')
    with pytest.raises(ValueError):
        calculate(path, 'sum', 'value')


def test_multiline_csv_requires_explicit_record_selection(tmp_path):
    path = tmp_path / 'multiline.csv'
    path.write_text('name,value\n"two\nlines",2\nnext,4\n', encoding='utf-8')
    with pytest.raises(ValueError, match='Multi-line CSV'):
        lookup(path, 3)
    with pytest.raises(ValueError, match='Multi-line CSV'):
        calculate(path, 'sum', 'value')


def test_summary_rows_do_not_double_count(tmp_path):
    path = tmp_path / 'totals.csv'
    path.write_text('name,value\na,2\nb,3\n合计,5\n', encoding='utf-8')
    with pytest.raises(ValueError, match='Summary rows'):
        calculate(path, 'sum', 'value')
    assert calculate(path, 'sum', 'value', filters={'name': 'a'})['value'] == 2


@pytest.fixture
def sensors(tmp_path):
    path=tmp_path/'sensors.xlsx'
    book=Workbook();ws=book.active;ws.title='传感器库存'
    ws.merge_cells('A1:D1');ws['A1']='库存采样记录'
    ws.append(['设备','温度','温度','压力(kPa)'])
    ws.append(['单位','（°C）','（°F）','（kPa）'])
    for number in range(4,11):ws.append([f'sensor-{number}',20+number,60+number,700+number])
    book.save(path);book.close()
    return path


def test_coordinate_header_lookup_uses_detected_row_not_anchor_value(inventory,sensors):
    import hashlib
    before={path:hashlib.sha256(path.read_bytes()).hexdigest() for path in (inventory,sensors)}
    note={}
    result=infer('库存管理附表3中，B2所在列的表头是什么？',inventory,diagnostics=note)
    assert result['operation']=='header_lookup' and result['value']=='原值'
    assert result['requested_cell']=='B2' and result['sheet']=='附表3-设备原值统计'
    assert result['row_start']==result['row_end']==result['header_row']==2
    assert result['column_letter']=='B' and result['column_index']==2
    assert result['range']=='B2:B2' and list(result['cells'])==['B2']
    assert result['unit']=='万元' and result['unit_source']=={'coordinate':'B3','row':3,'value':'（万元）'}
    assert note['reason']=='physical_header_lookup'
    later=infer('传感器库存中D10所处列的列名是什么？',sensors)
    assert later['value']=='压力(kPa)' and later['value']!=710
    assert later['requested_cell']=='D10' and later['range']=='D2:D2'
    assert later['cells']['D2']['unit_source']['coordinate']=='D3'
    assert all(hashlib.sha256(path.read_bytes()).hexdigest()==value for path,value in before.items())


def test_duplicate_header_labels_are_resolved_only_by_coordinate(sensors):
    left=infer('B10所在列的表头是什么',sensors)
    right=infer('C10所在列的表头是什么',sensors)
    assert left['value']==right['value']=='温度'
    assert left['range']=='B2:B2' and left['unit']=='°C'
    assert right['range']=='C2:C2' and right['unit']=='°F'
    with pytest.raises(ValueError,match='unique'):
        lookup(sensors,10,'温度')


def test_csv_coordinate_header_lookup_uses_physical_record_and_preserves_bytes(daily):
    before=daily.read_bytes()
    result=infer('daily.csv 中 B2 所在列的表头是什么',daily)
    assert result['value']=='kind' and result['range']=='B1:B1' and result['sheet'] is None
    assert result['requested_cell']=='B2' and result['unit'] is None
    assert infer('What is the column header of $D$4?',daily)['value']=='tp'
    assert daily.read_bytes()==before


def test_csv_duplicate_headers_and_units_use_column_identity(tmp_path):
    path=tmp_path/'units.csv';path.write_text('name,value,value\n单位,(kg),(m)\na,4,5\n',encoding='utf-8')
    assert infer('B3的表头是什么',path)['unit']=='kg'
    assert infer('C3的表头是什么',path)['unit']=='m'
    assert infer('C3的表头是什么',path)['range']=='C1:C1'


@pytest.mark.parametrize('coordinate',['E10','D11','A0','XFE2','ZZZ1','A1000021'])
def test_header_coordinates_outside_source_are_rejected(sensors,coordinate):
    with pytest.raises(ValueError,match='range|does not exist'):
        infer(coordinate+'所在列的表头是什么',sensors)


@pytest.mark.parametrize('question',[
    'B2和D10所在列的表头是什么','B2:B10的表头是什么','B2、B2的表头是什么',
    'B2所在列和C列的表头是什么','B2所在B:D列的表头是什么','B2所在列的第3列表头是什么',
    'B2所在列温度大于20的表头是什么','B2所在列的表头并求和','B2在2024年对应的表头是什么',
    'B2的表头 where kind=a','B2所处列中仅传感器的表头是什么','B2所处列按站点分组的表头是什么',
    'B2所在列中设备为传感器时的表头是什么',
])
def test_header_lookup_never_drops_multiple_coordinates_or_conditions(sensors,question):
    note={}
    assert infer(question,sensors,diagnostics=note) is None
    assert note['reason']=='needs_structured_filter'


def test_header_lookup_keeps_worksheet_disambiguation(inventory):
    assert infer('B2所在列的表头是什么',inventory) is None
    assert infer('附表3和附表13中B2的表头是什么',inventory) is None
    assert infer('附表13中B2的表头是什么',inventory)['range']=='B1:B1'
    assert infer('附表30中B2的表头是什么',inventory) is None


@pytest.mark.parametrize('merge',['B2:C2','B1:C1','B2:B6'])
def test_merged_header_hierarchy_is_explicitly_unsupported(sensors,merge):
    from openpyxl import load_workbook
    book=load_workbook(sensors);ws=book.active;ws.unmerge_cells('A1:D1')
    ws.merge_cells(merge);ws[merge.split(':')[0]]='分组标题'
    book.save(sensors);book.close()
    before=sensors.read_bytes()
    with pytest.raises(ValueError,match='merged or multi-row'):
        infer('B10所在列的表头是什么',sensors)
    assert sensors.read_bytes()==before


def test_short_vertical_header_merge_uses_original_top_cell_and_ignores_neighbors(tmp_path):
    path=tmp_path/'inventory-merged.xlsx';book=Workbook();ws=book.active;ws.title='库存附表7'
    ws.merge_cells('A1:E1');ws['A1']='虚构仓库设备清单'
    ws.append(['编号','装置名称','型号','台数','单价'])
    for column in ('A','B','C','D'):ws.merge_cells(f'{column}2:{column}3')
    ws['E3']='（元）';ws.append([1,'虚构流量计','X-1',2,10])
    book.save(path);book.close();before=path.read_bytes()
    result=infer('库存附表7中B2所在列的表头是什么？',path)
    assert result['value']=='装置名称' and result['range']=='B2:B2'
    assert result['header_merge_range']=='B2:B3' and result['cells']['B2']['merged_range']=='B2:B3'
    assert result['unit'] is None and result['requested_cell']=='B2'
    assert infer('B3所处列的列名是什么',path)['value']=='装置名称'
    assert infer('B4所处列的列名是什么',path)['value']=='装置名称'
    assert path.read_bytes()==before


def test_unrelated_horizontal_merged_headers_do_not_block_explicit_target(sensors):
    from openpyxl import load_workbook
    book=load_workbook(sensors);ws=book.active;ws.merge_cells('B2:C2');ws['B2']='温度分组'
    book.save(sensors);book.close();before=sensors.read_bytes()
    result=infer('D10所在列的表头是什么',sensors)
    assert result['value']=='压力(kPa)' and result['range']=='D2:D2'
    assert result['header_merge_range'] is None
    with pytest.raises(ValueError,match='requested column'):
        infer('C10所在列的表头是什么',sensors)
    assert sensors.read_bytes()==before


@pytest.mark.parametrize('address,value,error',[
    ('D2','=1+2','formula cache at D2'),('D10','=1+2','formula cache at D10'),
    ('D2','#DIV/0!','Excel error at D2'),('D10','#VALUE!','Excel error at D10'),
])
def test_header_lookup_preserves_formula_and_excel_error_guards(sensors,address,value,error):
    from openpyxl import load_workbook
    book=load_workbook(sensors);book.active[address]=value;book.save(sensors);book.close()
    with pytest.raises(ValueError,match=error):infer('D10所在列的表头是什么',sensors)


def test_blank_header_and_contradictory_units_are_not_guessed(sensors):
    from openpyxl import load_workbook
    book=load_workbook(sensors);book.active['D2']=None;book.save(sensors);book.close()
    with pytest.raises(ValueError,match='no explicit header'):infer('D10所在列的表头是什么',sensors)
    book=load_workbook(sensors);book.active['D2']='压力(Pa)';book.save(sensors);book.close()
    with pytest.raises(ValueError,match='Conflicting explicit units'):infer('D10所在列的表头是什么',sensors)


def test_header_lookup_does_not_treat_multiline_csv_as_one_physical_row(tmp_path):
    path=tmp_path/'multi.csv';path.write_text('name,value\n"two\nlines",2\n',encoding='utf-8')
    with pytest.raises(ValueError,match='Multi-line CSV'):infer('B2所在列的表头是什么',path)


@pytest.mark.parametrize('coordinate',[None,True,'B2:D3','B','2','BCDE2'])
def test_direct_header_lookup_requires_one_cell(sensors,coordinate):
    with pytest.raises(ValueError,match='exactly one'):header_lookup(sensors,coordinate)
