import json
from pathlib import Path
import pytest
from ragapp.store import Store
from ragapp.config import Settings
from ragapp.service import Service
from ragapp.api import create_app
from ragapp.calculator import calculate
from ragapp.parsers import parse
from ragapp.evaluation import grade

@pytest.fixture
def service(tmp_path):
    """创建隔离的 RAG 服务测试夹具。

    Args:
        tmp_path: Path：pytest 为当前测试提供的独立临时目录。

    Returns:
        Service：临时 papers/state 与 SQLite 台账；不启动 Qdrant 或调用模型。
    """
    corpus=tmp_path/'papers'; corpus.mkdir()
    return Service(Settings(root=tmp_path,corpus=corpus,state=tmp_path/'state'))

def test_delete_and_version_filter(service):
    """验证只有当前文档版本可召回，排除后所有块均不可召回。

    Args:
        service: Service：测试夹具创建的临时目录与 SQLite 台账，不依赖真实知识库。

    Returns:
        None：使用断言判断测试结果；不满足条件时测试失败。
    """
    with service.store.db() as db:
        db.execute("INSERT INTO documents VALUES('d','a.txt','h','v2','ready',NULL,0,0)")
        db.executemany('INSERT INTO chunks VALUES(?,?,?,?,?,?,?)',[(i,'d',v,'facts','{}','','') for i,v in [('old','v1'),('new','v2')]])
    assert set(service.store.active(['old','new']))=={'new'}
    service.store.exclude('d')
    assert not service.store.active(['new'])

def test_missing_root_does_not_delete(service):
    """验证资料根目录消失时停止同步，而非把全部文档认定为删除。

    Args:
        service: Service：测试夹具创建的临时目录与 SQLite 台账，不依赖真实知识库。

    Returns:
        None：使用断言判断测试结果；不满足条件时测试失败。
    """
    service.settings.corpus.rmdir()
    with pytest.raises(RuntimeError,match='deletions were NOT'): service.indexer.scan()

def test_scanner_excludes_outputs_and_secrets(service):
    """验证扫描只收集原始资料并排除 .env 与派生输出。

    Args:
        service: Service：测试夹具创建的临时目录与 SQLite 台账，不依赖真实知识库。

    Returns:
        None：使用断言判断测试结果；不满足条件时测试失败。
    """
    root=service.settings.corpus
    (root/'safe.md').write_text('hello'); (root/'.env').write_text('secret')
    (root/'extracted').mkdir(); (root/'extracted'/'output.md').write_text('derived')
    assert [p.name for p in service.indexer.scan()]==['safe.md']

def test_calculator_filters_and_no_code(tmp_path):
    """验证聚合与过滤数值正确，并拒绝非白名单操作。

    Args:
        tmp_path: Path：pytest 为当前测试提供的独立临时目录。

    Returns:
        None：使用断言判断测试结果；不满足条件时测试失败。
    """
    path=tmp_path/'a.csv'; path.write_text('kind,value\na,2\nb,4\na,8\n')
    assert calculate(path,'sum','value',filters={'kind':'a'})['value']==10
    assert calculate(path,'mean','value')['value']==pytest.approx(14/3)
    with pytest.raises(ValueError): calculate(path,'__import__','value')

def test_python_is_never_executed(tmp_path):
    """验证 Python 资料仅被 AST 解析，文件内抛异常的语句不会执行。

    Args:
        tmp_path: Path：pytest 为当前测试提供的独立临时目录。

    Returns:
        None：使用断言判断测试结果；不满足条件时测试失败。
    """
    path=tmp_path/'source.py'; path.write_text('raise RuntimeError("do not execute")\n\ndef f():\n    return 7\n')
    result=parse(path,tmp_path/'cache')
    assert any(r['locator'].get('symbol')=='f' and r['locator']['line_start']==3 for r in result)

def test_excel_title_units_and_missing_formula_cache(tmp_path):
    """验证表头/单位行识别和求和正确，并拒绝缺少缓存的公式。

    Args:
        tmp_path: Path：pytest 为当前测试提供的独立临时目录。

    Returns:
        None：使用断言判断测试结果；不满足条件时测试失败。
    """
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active;ws.append(['设备台账']);ws.append(['设备','原值']);ws.append([None,'（万元）']);ws.append(['A',2]);ws.append(['B',3])
    path=tmp_path/'inventory.xlsx';wb.save(path)
    assert calculate(path,'sum','原值')['value']==5
    result=parse(path,tmp_path/'cache')
    assert any('B2=原值' in item['text'] and '（万元）' in item['text'] for item in result)
    ws['B5']='=1+2';wb.save(path)
    with pytest.raises(ValueError,match='formula cache'):calculate(path,'sum','原值')

def test_docx_locators_have_no_fake_pages(tmp_path):
    """验证 Word 段落/表格定位有效且不虚构页码。

    Args:
        tmp_path: Path：pytest 为当前测试提供的独立临时目录。

    Returns:
        None：使用断言判断测试结果；不满足条件时测试失败。
    """
    from docx import Document
    doc=Document(); doc.add_heading('Header',1); doc.add_paragraph('Fact')
    table=doc.add_table(rows=1,cols=2); table.cell(0,0).text='A'; table.cell(0,1).text='2'
    path=tmp_path/'a.docx'; doc.save(path)
    result=parse(path,tmp_path/'cache')
    assert all('page' not in r['locator'] for r in result)
    assert any(r['locator'].get('table')==1 for r in result)

def test_api_blackbox_cannot_access_admin(service):
    """验证黑盒令牌不可访问原文或管理接口，无令牌检索返回未认证。

    Args:
        service: Service：测试夹具创建的临时目录与 SQLite 台账，不依赖真实知识库。

    Returns:
        None：使用断言判断测试结果；不满足条件时测试失败。
    """
    from fastapi.testclient import TestClient
    app=create_app(service,background_sync=False,ui=False)
    tokens=json.loads((service.settings.state/'tokens.json').read_text())
    client=TestClient(app); headers={'Authorization':'Bearer '+tokens['blackbox']}
    for method,url in [('get','/api/documents'),('get','/api/documents/a/source'),('post','/api/sync'),('delete','/api/documents/a'),('get','/api/costs')]:
        assert getattr(client,method)(url,headers=headers).status_code==403
    assert client.post('/api/search',json={'question':'hello'}).status_code==401

def test_grading_cannot_pass_without_facts_and_citations():
    """验证答案只有正确数值但缺少引用时无法通过辅助评分。

    Args:
        无显式输入参数。

    Returns:
        None：使用断言判断测试结果；不满足条件时测试失败。
    """
    case={'answerable':True,'facts':[['42']],'evidence':[]}
    assert not grade(case,{'answer':'42','citations':[]})['pass']

def test_paid_budget_blocks_before_network(service,monkeypatch):
    """验证预算耗尽时在发起付费网络请求前阻止调用。

    Args:
        service: Service：测试夹具创建的临时目录与 SQLite 台账，不依赖真实知识库。
        monkeypatch: pytest.MonkeyPatch：替换网络或其他外部依赖，测试结束后自动恢复。

    Returns:
        None：使用断言判断测试结果；不满足条件时测试失败。
    """
    service.models.keys={'DEEPSEEK_API_KEY':'fake'}
    with service.store.db() as db:
        db.execute("INSERT INTO calls VALUES('spent','production','development','m','m','complete',30,30,'{}',NULL,0,0)")
    def network(*a,**k):
        """把任何意外发出的网络调用立即判为测试失败。

        Args:
            a: tuple：兼容被替换 httpx.post 的位置参数。
            k: dict：兼容被替换 httpx.post 的关键字参数。

        Returns:
            None：调用时通过 pytest.fail 抛出失败异常，不正常返回。
        """
        pytest.fail('Must not make a paid request after budget limit')
    monkeypatch.setattr('httpx.post',network)
    with pytest.raises(RuntimeError,match='budget'): service.models.complete([{'role':'user','content':'Hi'}])
