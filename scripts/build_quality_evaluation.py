"""Build a reproducible, source-derived 120/120 evaluation DRAFT, never from RAG output.

Gold text is read from original PDF/DOCX/text/table files. Visual gold records
direct agent inspection separately from human review. A frozen draft is not an
independently audited benchmark; every question retains that distinction.
"""
import argparse
import csv
import difflib
import hashlib
import io
import json
import re
import sqlite3
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from ragapp.evaluation import normalized, save_report
from ragapp.parsers import read_text


def sha(value):
    return hashlib.sha256(value if isinstance(value, bytes) else value.encode()).hexdigest()


def source_blocks(path):
    """Extract directly without invoking OCR, vision providers, or the RAG index."""
    ext = path.suffix.lower()
    if ext == '.pdf':
        from pypdf import PdfReader
        return [{'text': p.extract_text() or '', 'locator': {'page': i}, 'extraction': 'original_pdf_pypdf'}
                for i, p in enumerate(PdfReader(path).pages, 1)]
    if ext == '.docx':
        from docx import Document
        doc = Document(path)
        blocks = [{'text': p.text, 'locator': {'paragraph': i}, 'extraction': 'original_docx_xml'}
                  for i, p in enumerate(doc.paragraphs, 1)]
        blocks += [{'text': ' | '.join(c.text for c in row.cells), 'locator': {'table': i, 'row': j},
                    'extraction': 'original_docx_xml'} for i, table in enumerate(doc.tables, 1)
                   for j, row in enumerate(table.rows, 1)]
        return blocks
    if ext == '.xlsx':
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        wb = load_workbook(path, read_only=True, data_only=True)
        blocks = []
        try:
            for sheet in wb:
                # Avoid personnel/contact tables in generated text questions.
                if re.search(r'人员|人才|通讯|联系', sheet.title):
                    continue
                for index, row in enumerate(sheet.iter_rows(), 1):
                    values = [f'{c.coordinate}={c.value}' for c in row if c.value is not None]
                    if values:
                        blocks.append({'text': ' | '.join(values), 'locator': {'sheet': sheet.title,
                            'range': f'A{index}:{get_column_letter(len(row))}{index}'}, 'extraction': 'original_xlsx_values'})
        finally:
            wb.close()
        return blocks
    if ext == '.csv':
        rows = csv.DictReader(io.StringIO(read_text(path)))
        return [{'text': ' | '.join(f'{k}={v}' for k, v in row.items()), 'values': row,
                 'locator': {'row': i}, 'extraction': 'original_csv'} for i, row in enumerate(rows, 2)]
    if ext in {'.png', '.jpg', '.jpeg', '.webp', '.bmp'}:
        return []
    lines = read_text(path).splitlines()
    return [{'text': line, 'locator': {'line': i}, 'extraction': 'original_text'}
            for i, line in enumerate(lines, 1) if line.strip()]


def topic(path):
    name = Path(path).stem
    if '基于LSTM-Seq2Seq' in name:
        return 'LSTM-Seq2Seq流域径流预测研究'
    if name == 'lstm':
        return '课程目录的径流预测程序' if '/' in path else '根目录的径流预测程序'
    if name == 'readme':
        return '径流预测课程项目的作者说明'
    if name == 'requirements':
        return '径流预测项目依赖清单'
    if 'IPCC' in name:
        return 'IPCC AR6 WGII水资源章节'
    if name == 'ERA5_discharge_snow':
        return 'ERA5日度气象径流记录'
    if '附表1-12' in name:
        return '自然灾害监测预警工程研究中心附表'
    return re.sub(r'^\d+\.', '', name)


def domain(path):
    if re.search(r'LSTM|lstm|ERA5|IPCC|水库|经济成本|框图|readme|requirements', path):
        return 'water'
    if re.search(r'CCD|TCD|STM32|检测报告', path):
        return 'electronics'
    if re.search(r'机器学习|统计分类|线性模型|SVM|神经网络|聚类|降维|贝叶斯', path):
        return 'learning'
    return 'other'


def family(path):
    # Exports of the same report are one split family to avoid cross-format leaks.
    if '基于LSTM-Seq2Seq' in Path(path).stem:
        return 'runoff-report-exports'
    return Path(path).stem


def make_evidence(doc, block, quote):
    return {'path': doc['path'], 'source_hash': doc['hash'], 'locator': block['locator'],
            'span_text': quote, 'span_sha256': sha(normalized(quote)), 'extraction': block['extraction']}


def usable_sentences(block):
    text = block['text'].strip()
    if not text or re.search(r'MIRROR|image contains|no OCR|免责声明|冒充百度|版权|copyright', text, re.I):
        return []
    # Keep original characters and linebreaks in quotations.
    sentences = re.split(r'(?<=[。！？;；])\s*|\n(?=[•]|\d+[、.)])', text)
    result = []
    for sentence in sentences:
        sentence = sentence.strip(' \n\t•')
        compact = normalized(sentence)
        letters = len(re.findall(r'[\u4e00-\u9fffA-Za-z]', sentence))
        if not 28 <= len(compact) <= 230 or letters < 18 or letters / len(compact) < .60:
            continue
        if re.search(r'[\ufffd\ue000-\uf8ff]|[?？]|By\s+机器学习|\.{4,}|…{2,}|[─━]{3,}', sentence):
            continue
        if re.search(r'^\s*(参考文献|References|\[\d+\]|https?://)|^\d+(?:\.\d+){2,}\s', sentence):
            continue
        result.append(sentence)
    return result


def question_candidate(doc, block, quote, split):
    compact = re.sub(r'\s+', ' ', quote).strip(' *#•-')
    # An anchor is a subject/prefix, never the target fact itself.
    match = next((m for m in re.finditer(r'定义为|包括|包含|采用|用于|需要|表示|可以|能够|支持|称为|等于|意味着|说明|表明|反映|导致|取决于|是|：|:', compact)
                  if 4 <= m.start() <= 70), None)
    if match is None:
        match = next((m for m in re.finditer(r'\b(?:is|are|can|include|includes|provides|require|requires|shows|indicates)\b', compact, re.I)
                      if 6 <= m.start() <= 70), None)
    if match is None:
        return None
    cut = match.start()
    anchor = compact[:cut].strip(' *#，,：:')
    remainder = compact[cut:].strip(' *#，,：:')
    if len(anchor) < 5 or len(remainder) < 12:
        return None
    stem = topic(doc['path'])
    if split == 'dev':
        question = f'根据{stem}的资料，关于“{anchor}”，原文给出的具体说明是什么？'
        template = 'subject-explanation-dev'
    else:
        question = f'在{stem}这一背景下，请找出“{anchor}”所对应的后续结论或取值，并给出依据。'
        template = 'continuation-evidence-blind'
    return {'question': question, 'evidence': [make_evidence(doc, block, quote)],
            'facts': [[re.escape(remainder[:min(22, len(remainder))])]], 'answer_hint': remainder,
            'topic': stem, 'domain': domain(doc['path']), 'template_group': template}


VISUAL_CASES = [
    ('107541666_p0.png', '资料中的粉色兽耳插画，人物头发与背景心形图案主要是什么颜色？',
     ['粉色', '粉色与紫色'], 'Pink-haired illustrated character; background hearts are pink and purple.'),
    ('a4037845571a4e9db6833c9dfcfe7c52~tplv-p14lwwcsbr-origin-webp_q75.webp',
     '城市背景中的银色面部角色图片，胸口中央的装置是什么形状、什么颜色？',
     ['圆形', '蓝色'], 'The center of the chest has a round pale-blue device.'),
    ('大学一班.jpg', '毕业合影背景板上的两句大字标语是什么？',
     ['青春不散场', '梦想正启航'], '青春不散场，梦想正启航。'),
    ('论文框图.jpg', '水库研究框图中，配置优化的三个目标是什么？',
     ['供水效益', '总缺水量', 'Gini'], '供水效益最大；总缺水量最小；满足率Gini最小。'),
    ('peak0.6 w5 180d loss.png', '课程项目的损失曲线图同时展示了哪两组数据的损失？',
     ['训练', '验证'], '训练集 Loss 与验证集 Loss。'),
    ('180d 0.6peak w5 1-200-399.png', '课程项目的多子图径流预测图中，图例对照的两类序列是什么？',
     ['真实', '预测'], '真实月均径流 Ground Truth 与预测月均径流 Prediction。'),
]

NEGATIVE_QUESTIONS = {
    'dev': [
        '当前资料是否给出了径流预测模型在真实生产环境连续运行一年的宕机次数？',
        '当前资料是否记录了LSTM训练在英伟达H100上的实测总耗电量？',
        '能否根据现有项目资料给出每个预测结果经过独立第三方盲测的原始评分？',
        '现有资料有没有说明径流模型公开部署服务器的公网地址与正式上线日期？',
        '当前资料是否提供了TCD1304驱动板连续老化一万小时后的完整故障日志？',
        '这些CCD资料有没有提供全部制造批次的逐器件出厂校准证书？',
        '电机监测论文是否给出了监测系统五年质保期间实际赔付记录？',
        '已有课件是否记录了本届每一位学生对机器学习课程的匿名满意度原始问卷？',
        '资料是否提供了水库多目标优化软件经过正式认证的安全等级证书编号？',
        '能否从现有RAG介绍资料确定本项目检索服务在线运行的真实月活跃用户数？',
        '依赖清单能否证明这些库在项目作者电脑上安装成功的逐包安装日志？',
        '作者说明是否提供了每段AI生成代码对应的完整模型会话及逐行人工审核记录？',
    ],
    'blind': [
        '请给出ERA5日度气象径流记录在2045年1月1日的实测discharge值。',
        '现有资料能否给出2040年努列克水库每天的真实入库径流观测？',
        '只依据资料，2038年哈特隆州实际农业灌溉用水总量是多少？',
        '资料中的电机在2035年实际发生了多少次轴承故障？',
        '这套CCD驱动板在2032年的真实市场销售收入是多少？',
        'IPCC水资源章节能否给出未来2035年每一次洪水的确切发生日期与损失？',
        '根据已上传资料，下一次训练运行最后一个epoch的实际验证损失会是多少？',
        '资料中的毕业合影能够证明每位同学毕业十年后的就业单位和工资吗？',
        '资料中的插画是否附有创作者已签署的商业授权合同与授权期限？',
        '城市背景角色图片是否记录了准确拍摄地点的GPS坐标与拍摄时间戳？',
        '拟录取名单是否提供了所有考生毕业后的真实学位授予记录？',
        '当前设备附表能否证明表中设备下个月的真实采购中标价格？',
    ],
}


def build(root, output):
    output.mkdir(parents=True, exist_ok=True)
    if (output / 'manifest.json').exists():
        raise ValueError('Frozen dataset already exists; use a new directory rather than overwrite gold')
    db = sqlite3.connect((root / '.rag/catalog.sqlite3').resolve().as_uri() + '?mode=ro', uri=True)
    db.row_factory = sqlite3.Row
    docs = [dict(d) for d in db.execute("SELECT * FROM documents WHERE excluded=0 AND status IN ('ready','update_failed') ORDER BY path")]
    db.close()
    pool_path = output / 'original-source-pool.json'
    original_pool = json.loads(pool_path.read_text(encoding='utf-8')) if pool_path.exists() else {}
    for doc in docs:
        doc['path'] = doc['path'].replace('\\', '/')
        source = root / 'papers' / doc['path']
        if sha(source.read_bytes()) != doc['hash']:
            raise ValueError('Source changed before gold construction: ' + doc['path'])
        doc['blocks'] = original_pool.get(doc['hash']) or source_blocks(source)
        original_pool[doc['hash']] = doc['blocks']
    save_report(pool_path, original_pool)
    sources = {d['path']: {'hash': d['hash'], 'version_at_authoring': d['version'], 'id': d['id']} for d in docs}
    families = sorted({family(d['path']) for d in docs})
    assignments = {f: ('dev' if i % 2 == 0 else 'blind') for i, f in enumerate(families)}
    pools = {'dev': [], 'blind': []}
    used_spans = set()
    for doc in docs:
        split = assignments[family(doc['path'])]
        stem = Path(doc['path']).stem
        for block in doc['blocks']:
            if stem == 'readme':
                question = ('作者说明如何区分径流模型构建与测试代码中的AI参与？' if 'AI生成' in block['text']
                            else '作者说明建议优先查看课程报告的哪些文件格式？')
                pools[split].append({'question': question,
                    'evidence': [make_evidence(doc, block, block['text'].strip())],
                    'facts': [['测试的代码', 'PDF版本']], 'answer_hint': block['text'].strip(),
                    'template_group': 'authorship-and-format-' + split, 'topic': topic(doc['path']), 'domain': 'water'})
                continue
            if '拟录取名单' in stem:
                for line in block['text'].splitlines():
                    match = re.match(r'^(\d{15})\s+(\S+)\s+(.*)', line.strip())
                    if match and '全日制' in line:
                        identifier, name, rest = match.groups()
                        scores = re.findall(r'\b\d+\.\d{2}\b', rest)
                        if scores:
                            pools[split].append({'question': f'2024年拟录取名单中，考生编号{identifier}对应的加权总成绩是多少？',
                                'evidence': [make_evidence(doc, block, line.strip())], 'facts': [[re.escape(scores[-1])]],
                                'answer_hint': scores[-1], 'template_group': 'admission-weighted-score-' + split,
                                'topic': topic(doc['path']), 'domain': 'other'})
            if stem == '机器学习作业':
                match = re.search(r'截距\s*w0\s*=\s*0\.3667\s*斜率\s*w1\s*=\s*0\.4000', block['text'])
                if match:
                    pools[split].append({'question': '第四章线性回归习题给出的截距与斜率分别是多少？请引用PDF习题结果。',
                        'evidence': [make_evidence(doc, block, match.group())], 'facts': [['0\\.3667'], ['0\\.4000']],
                        'answer_hint': '截距0.3667，斜率0.4000', 'template_group': 'regression-output-' + split,
                        'topic': topic(doc['path']), 'domain': 'learning'})
                elif re.search(r'斜率\s*w1\s*=\s*0\.4000', block['text']):
                    quote = block['text'].strip()
                    pools[split].append({'question': '第四章线性回归习题在Markdown解答中报告的斜率是多少？',
                        'evidence': [make_evidence(doc, block, quote)], 'facts': [['0\\.4000']],
                        'answer_hint': '0.4000', 'template_group': 'regression-slope-' + split,
                        'topic': topic(doc['path']), 'domain': 'learning'})
            if doc['path'].endswith('.xlsx'):
                match = re.search(r'(?:^| \| )(B\d+)=([^|]+)', block['text'])
                if match and len(match.group(2).strip()) >= 4 and not re.search(r'姓名|序号|项目名称|设备名称', match.group(2)):
                    cell, value = match.groups()
                    quote = f'{cell}={value.strip()}'
                    pools[split].append({'question': f'工程研究中心的{block["locator"]["sheet"]}中，{cell}单元格记载的项目或内容是什么？',
                        'evidence': [make_evidence(doc, block, quote)], 'facts': [[re.escape(value.strip())]],
                        'answer_hint': value.strip(), 'template_group': 'spreadsheet-cell-' + split,
                        'topic': topic(doc['path']), 'domain': 'other'})
            # Requirements and code assignments get concrete factual questions.
            match = re.fullmatch(r'([A-Za-z0-9_]+)==([\w.]+)', block['text'].strip()) if stem == 'requirements' else None
            if match:
                name, value = match.groups()
                quote = block['text'].strip()
                candidate = {'question': f'径流预测项目的依赖清单将{name}固定在什么版本？',
                    'evidence': [make_evidence(doc, block, quote)], 'facts': [[re.escape(value)]],
                    'answer_hint': value, 'template_group': 'dependency-version-' + split,
                    'topic': topic(doc['path']), 'domain': domain(doc['path'])}
                pools[split].append(candidate)
                continue
            if stem == 'lstm':
                match = re.match(r'^\s*([A-Z][A-Z_]{3,})\s*=\s*([^#\n]+)', block['text'])
                if match:
                    symbol, value = match.groups()
                    candidate = {'question': f'{topic(doc["path"])}中，{symbol}配置为什么值？',
                        'evidence': [make_evidence(doc, block, match.group().strip())],
                        'facts': [[re.escape(value.strip())]], 'answer_hint': value.strip(),
                        'template_group': 'code-constant-' + symbol, 'topic': topic(doc['path']), 'domain': 'water'}
                    # The same symbol's variants belong together, even across files.
                    pools['dev' if int(sha(symbol), 16) % 2 == 0 else 'blind'].append(candidate)
                continue
            if block.get('values'):
                number = block['locator']['row']
                if number % 347 != 2:
                    continue
                values = block['values']
                for column in ('discharge', 'tp'):
                    date, value = values.get('date'), values.get(column)
                    if date and value:
                        candidate = {'question': f'ERA5气象径流记录中，{date}的{column}原始记录值是多少？',
                            'evidence': [make_evidence(doc, block, f'{column}={value}')],
                            'facts': [[re.escape(value)]], 'answer_hint': value,
                            'template_group': 'csv-value-' + split, 'topic': topic(doc['path']), 'domain': 'water'}
                        pools[split].append(candidate)
                continue
            for quote in usable_sentences(block):
                key = sha(re.sub(r'[^\w]', '', normalized(quote)))
                if key in used_spans:
                    continue
                used_spans.add(key)
                candidate = question_candidate(doc, block, quote, split)
                if candidate:
                    pools[split].append(candidate)
    cases, source_checks = [], []
    taken = set()
    def add(candidate, split, category, serial):
        case = {**candidate, 'id': f'qv2-{split}-{category}-{serial:03d}', 'schema_version': 'quality-v2',
            'split': split, 'category': category, 'answerable': True, 'review_status': 'pending_human_review',
            'authoring': 'source-derived deterministic draft; no optimized-RAG outputs',
            'cluster_id': sha('|'.join(sorted(e.get('span_sha256', e['path']) for e in candidate.get('evidence', []))))[:20]}
        cases.append(case)
        for evidence in case.get('evidence', []):
            taken.add((evidence['path'], json.dumps(evidence['locator'], sort_keys=True), evidence.get('span_sha256')))
            source_checks.append({'case_id': case['id'], **evidence,
                'original_hash_verified': True, 'span_extracted_from_original': bool(evidence.get('span_text')),
                'human_review': 'pending'})
        return case
    def unused(item):
        return all((e['path'], json.dumps(e['locator'], sort_keys=True), e.get('span_sha256')) not in taken for e in item['evidence'])
    # One visual case per original image. Region observations are not fake text spans.
    for name, question, facts, observation in VISUAL_CASES:
        doc = next((d for d in docs if Path(d['path']).name == name), None)
        if doc:
            split = assignments[family(doc['path'])]
            item = {'question': question, 'facts': [[re.escape(f)] for f in facts], 'scorable': False,
                'visual_observation': observation, 'visual_review': 'agent_original_image_inspection; independent human pending',
                'template_group': 'visual-' + split, 'evidence': [{'path': doc['path'], 'source_hash': doc['hash'],
                    'locator': {'image': name}, 'extraction': 'original_image_visual_inspection'}]}
            add(item, split, 'single', sum(c['split'] == split and c['category'] == 'single' for c in cases) + 1)
    # Seed coverage before filling any large source's quota.
    for doc in docs:
        if any(e['path'] == doc['path'] for c in cases for e in c.get('evidence', [])):
            continue
        split = assignments[family(doc['path'])]
        candidates = [c for c in pools[split] if c['evidence'][0]['path'] == doc['path'] and unused(c)]
        if candidates:
            add(candidates[0], split, 'single', sum(c['split'] == split and c['category'] == 'single' for c in cases) + 1)
        else:
            raise ValueError('No source-grounded candidate for ' + doc['path'])
    for split in ('dev', 'blind'):
        pool = pools[split]
        # Pair two distinct sources in the same research domain.
        for index in range(12):
            available = [c for c in pool if unused(c)]
            first = available[(index * 19) % len(available)]
            second = next((c for c in available if c['domain'] == first['domain'] and
                           c['evidence'][0]['path'] != first['evidence'][0]['path']), None)
            if second is None:
                first = next(c for c in available if c['domain'] == 'learning')
                second = next(c for c in available if c['domain'] == 'learning' and c['evidence'][0]['path'] != first['evidence'][0]['path'])
            item = {'question': f'请分别依据两份资料回答并逐项引用：{first["question"]}另外，{second["question"]}',
                'facts': first['facts'] + second['facts'], 'evidence': first['evidence'] + second['evidence'],
                'answer_hint': [first['answer_hint'], second['answer_hint']], 'template_group': 'two-source-' + split}
            add(item, split, 'cross_source', index + 1)
        for index in range(12):
            available = [c for c in pool if unused(c)]
            candidate = available[(index * 23) % len(available)]
            item = {**candidate, 'history': [{'role': 'user', 'content': candidate['question']}],
                'question': '请接着刚才的问题给出具体的原文依据和结论，不要仅重复主题名称。',
                'template_group': 'followup-evidence-' + split}
            add(item, split, 'multi_turn', index + 1)
        # Round-robin across sources prevents the 8,516-row CSV dominating gold.
        queues = defaultdict(list)
        for item in pool:
            if unused(item):
                queues[item['evidence'][0]['path']].append(item)
        count = sum(c['split'] == split and c['category'] == 'single' for c in cases)
        while count < 84:
            progress = False
            for path in sorted(queues):
                while queues[path] and not unused(queues[path][0]):
                    queues[path].pop(0)
                if queues[path] and count < 84:
                    add(queues[path].pop(0), split, 'single', count + 1)
                    count += 1
                    progress = True
            if not progress:
                raise ValueError('Insufficient distinct source spans')
        for index, question in enumerate(NEGATIVE_QUESTIONS[split], 1):
            cases.append({'id': f'qv2-{split}-unanswerable-{index:03d}', 'schema_version': 'quality-v2',
                'split': split, 'category': 'unanswerable', 'question': question, 'answerable': False,
                'facts': [], 'evidence': [], 'review_status': 'pending_human_absence_review',
                'negative_basis': 'Requested operational/future/authorization records are not supplied; independently verify absence before acceptance.',
                'template_group': ('missing-operational-records' if split == 'dev' else 'unobserved-future-or-external-records'),
                'cluster_id': sha(question)[:20]})
    # Shared exact spans and template families cannot straddle splits.
    spans = defaultdict(set)
    templates = defaultdict(set)
    for case in cases:
        templates[case['template_group']].add(case['split'])
        for evidence in case.get('evidence', []):
            if evidence.get('span_sha256'):
                spans[evidence['span_sha256']].add(case['split'])
    leaked = [key for key, values in spans.items() if len(values) > 1]
    if leaked or any(len(v) > 1 for v in templates.values()):
        raise ValueError('Evidence-span or template-family split leakage')
    ordered = sorted(cases, key=lambda c: c['id'])
    datasets = {}
    for split in ('dev', 'blind'):
        subset = [c for c in ordered if c['split'] == split]
        if len(subset) != 120:
            raise ValueError('Expected exactly 120 cases per split')
        path = output / (split + '.json')
        save_report(path, subset)
        datasets[path.name] = {'sha256': sha(path.read_bytes()), 'count': len(subset)}
        save_report(output / (split + '-questions.json'), [{k: c[k] for k in ('id', 'question', 'history') if k in c} for c in subset])
    save_report(output / 'source-checks.json', source_checks)
    coverage = Counter(e['path'] for c in cases for e in c.get('evidence', []))
    names = [Path(d['path']).name.lower() for d in docs]
    no_filename = sum(not any(name in c['question'].lower() for name in names) for c in cases)
    manifest = {'schema_version': 'quality-v2', 'draft_revision': 3, 'status': 'frozen_draft_awaiting_independent_human_review',
        'sources': sources, 'datasets': datasets, 'counts': dict(Counter(c['category'] for c in cases)),
        'source_coverage': dict(coverage), 'covered_documents': len(coverage), 'questions_without_filename': no_filename,
        'human_verified_cases': 0, 'source_checks_sha256': sha((output / 'source-checks.json').read_bytes()),
        'split_policy': 'whole source families where possible; identical spans and code-symbol/template groups cannot cross splits',
        'leakage_check': {'exact_span_overlap': 0, 'template_group_overlap': 0,
                          'semantic_near_duplicate_review': 'pending_human_review'},
        'limitations': ['Deterministic source-completion wording needs editorial review before final blind acceptance.',
                       'Visual cases have original-image observations, no fabricated text spans; excluded from text-span metrics.',
                       'Negative-case absence and cross-source question usefulness require independent human review.',
                       'The blind label describes split isolation, not an OS access-control guarantee.']}
    save_report(output / 'manifest.json', manifest)
    (output / 'README.md').write_text('# Quality-v2 源材料评测草案\n\n'
        '共240题，开发/保留各120；跨来源、多轮、不可回答各24题，覆盖全部38份原资料。'
        '题目从原文件提取，未使用待优化检索结果或模型答案作为金标。\n\n'
        '**冻结不等于人工验收通过。** 当前全部题目仍需独立人工核对；源码/表格题和按主题检索原文的题目含模板化措辞，'
        '正式盲测前应审核语义、近重复、答案覆盖和不可回答依据。修改后必须发布新版本，不得覆盖已冻结题集。'
        '纯视觉题保留原图观察，不计入文本span召回。\n\n'
        'dev.json用于调参；blind.json仅在配置冻结后使用。source-checks.json给出原文摘录、SHA256和物理定位。'
        'manifest.json区分源文件校验、程序构造与人工审核状态。\n', encoding='utf-8')
    return manifest


def refresh(root, output, *, require_agent_reviews=True):
    """Re-freeze edited cases in place, never regenerate or reassign questions.

All validations finish before writing. The edited files plus preceding metadata
are archived first. Agent review remains distinct from independent human review.
"""
    root, output = Path(root), Path(output)
    input_hashes = {name: sha((output / name).read_bytes())
                    for name in ('dev.json', 'blind.json', 'manifest.json')}
    old_manifest = json.loads((output / 'manifest.json').read_text(encoding='utf-8'))
    datasets = {split: json.loads((output / (split + '.json')).read_text(encoding='utf-8'))
                for split in ('dev', 'blind')}
    sources = old_manifest['sources']
    pool = json.loads((output / 'original-source-pool.json').read_text(encoding='utf-8'))
    for path, source in sources.items():
        original = (root / 'papers' / path).resolve()
        if not original.is_relative_to((root / 'papers').resolve()) or not original.is_file():
            raise ValueError('Frozen source missing or outside corpus: ' + path)
        if sha(original.read_bytes()) != source['hash']:
            raise ValueError('Frozen original changed: ' + path)
    all_cases = datasets['dev'] + datasets['blind']
    if len({c['id'] for c in all_cases}) != 240:
        raise ValueError('Expected 240 unique case IDs')
    spans, clusters, templates = defaultdict(set), defaultdict(set), defaultdict(set)
    source_checks, coverage = [], Counter()
    machine_checked = visual_checked = 0
    for split, cases in datasets.items():
        if len(cases) != 120:
            raise ValueError('Each split must contain exactly 120 cases')
        counts = Counter(c['category'] for c in cases)
        if counts != Counter(single=84, cross_source=12, multi_turn=12, unanswerable=12):
            raise ValueError('Each split must retain 84/12/12/12 category quotas')
        for case in cases:
            if case.get('split') != split or case.get('schema_version') != 'quality-v2':
                raise ValueError('Case split or schema changed: ' + case['id'])
            review = case.get('agent_review') or {}
            if require_agent_reviews and (case.get('agent_source_reviewed') is not True
                    or review.get('source_checked') is not True or not review.get('reviewer') or not review.get('reason')):
                raise ValueError('Missing independent agent source review: ' + case['id'])
            if case.get('review_status') == 'human_verified':
                raise ValueError('Refresh cannot promote agent reviews to human verification')
            if not case.get('question', '').strip() or not case.get('cluster_id') or not case.get('template_group'):
                raise ValueError('Question, cluster and template group are required: ' + case['id'])
            clusters[case['cluster_id']].add(split)
            templates[case['template_group']].add(split)
            if case['category'] == 'multi_turn' and not any(m.get('role') == 'user' and m.get('content') for m in case.get('history', [])):
                raise ValueError('Multi-turn case lost its user history: ' + case['id'])
            if case['category'] == 'unanswerable':
                if case.get('answerable') is not False or case.get('evidence'):
                    raise ValueError('Unanswerable case has inconsistent gold: ' + case['id'])
                source_checks.append({'case_id': case['id'], 'answerable': False,
                    'absence_review': case.get('absence_review') or review.get('reason'),
                    'agent_source_reviewed': case.get('agent_source_reviewed', False),
                    'agent_review': review, 'human_review': 'pending'})
                continue
            if case.get('answerable') is not True or not case.get('evidence'):
                raise ValueError('Answerable case needs original evidence: ' + case['id'])
            if case['category'] == 'cross_source' and len({e['path'] for e in case['evidence']}) < 2:
                raise ValueError('Cross-source case requires distinct original sources: ' + case['id'])
            for evidence in case['evidence']:
                name = evidence['path'].replace('\\', '/')
                if name not in sources or evidence.get('source_hash') != sources[name]['hash']:
                    raise ValueError('Evidence source hash differs from frozen originals: ' + case['id'])
                evidence['path'] = name
                coverage[name] += 1
                quote = evidence.get('span_text')
                matched = False
                if quote:
                    evidence['span_sha256'] = sha(normalized(quote))
                    spans[evidence['span_sha256']].add(split)
                    expected = evidence.get('locator', {})
                    # Reviewed text spans can legitimately cover several original lines.
                    if 'line_start' in expected and Path(name).suffix.lower() in {'.md', '.txt', '.py'}:
                        lines = (root / 'papers' / name).read_text(encoding='utf-8-sig').splitlines()
                        lo, hi = expected['line_start'], expected.get('line_end', expected['line_start'])
                        matched = 1 <= lo <= hi <= len(lines) and normalized(quote) in normalized('\n'.join(lines[lo-1:hi]))
                    # The authoring pool is independent original extraction, not retrieval text.
                    for block in pool.get(sources[name]['hash'], []):
                        location = block['locator']
                        same = all(location.get(k) == v for k, v in expected.items()
                                   if k not in {'char_start', 'char_end', 'document_char_start', 'document_char_end'})
                        if same and normalized(quote) in normalized(block['text']):
                            matched = True
                            break
                    visual = review.get('visual_checked') is True or evidence.get('verification') == 'original_image_visual_inspection'
                    if not matched and not visual:
                        raise ValueError('Gold span does not occur at its original location: ' + case['id'])
                    machine_checked += matched
                    visual_checked += not matched and visual
                elif case.get('scorable') is not False:
                    raise ValueError('Text-scorable case has no span: ' + case['id'])
                source_checks.append({'case_id': case['id'], **evidence,
                    'original_hash_verified': True, 'span_extracted_from_original': matched,
                    'agent_source_reviewed': case.get('agent_source_reviewed', False),
                    'agent_review': review, 'human_review': 'pending'})
    if len(coverage) != len(sources):
        raise ValueError('Reviewed cases no longer cover every original source')
    if any(len(v) > 1 for mapping in (spans, clusters, templates) for v in mapping.values()):
        raise ValueError('Exact evidence span, cluster or template group leaks across splits')
    filenames = [Path(name).name.lower() for name in sources]
    no_filename = sum(not any(name in (c['question'] + ' '.join(m.get('content', '') for m in c.get('history', []))).lower()
                             for name in filenames) for c in all_cases)
    if no_filename < 120:
        raise ValueError('At least half the questions must omit filenames including their history')
    def query_text(case):
        return normalized(' '.join(m.get('content', '') for m in case.get('history', [])) + case['question'])
    near_pairs = []
    for left in datasets['dev']:
        a = query_text(left)
        for right in datasets['blind']:
            b = query_text(right)
            ratio = difflib.SequenceMatcher(None, a, b, autojunk=False).ratio()
            if min(len(a), len(b)) >= 30 and ratio >= .92:
                near_pairs.append({'dev_id': left['id'], 'blind_id': right['id'], 'similarity': ratio})
    if near_pairs:
        save_report(output / 'pending-leakage-review.json', near_pairs)
        raise ValueError('Near-duplicate cross-split questions require review; see pending-leakage-review.json')
    if any(sha((output / name).read_bytes()) != value for name, value in input_hashes.items()):
        raise ValueError('Dataset was edited during validation; retry after its authors finish')
    archive = output / 'review-refresh-history' / (str(time.time_ns()) + '-' + sha((output / 'manifest.json').read_bytes())[:12])
    archive.mkdir(parents=True)
    for name in ('dev.json', 'blind.json', 'dev-questions.json', 'blind-questions.json', 'manifest.json',
                 'source-checks.json', 'summary.json', 'README.md', 'dev-gold-audit.json', 'blind-gold-audit.json'):
        path = output / name
        if path.exists():
            (archive / name).write_bytes(path.read_bytes())
    dataset_hashes = {}
    for split, cases in datasets.items():
        path = output / (split + '.json')
        save_report(path, cases)
        dataset_hashes[path.name] = {'sha256': sha(path.read_bytes()), 'count': len(cases)}
        save_report(output / (split + '-questions.json'), [{k: c[k] for k in ('id', 'question', 'history') if k in c} for c in cases])
        audit_path = output / (split + '-gold-audit.json')
        if audit_path.exists():
            audit = json.loads(audit_path.read_text(encoding='utf-8'))
            audit['pre_refresh_dataset_sha256'] = audit.get('dataset_sha256')
            audit['dataset_sha256'] = dataset_hashes[path.name]['sha256']
            audit['refresh_note'] = 'Only evidence path/span hash metadata was normalized after independent source review.'
            save_report(audit_path, audit)
    save_report(output / 'source-checks.json', source_checks)
    summary = {'total': 240, 'splits': {'dev': 120, 'blind': 120}, 'covered_documents': len(coverage),
        'counts': dict(Counter(c['category'] for c in all_cases)), 'questions_without_filename': no_filename,
        'agent_source_reviewed_cases': sum(c.get('agent_source_reviewed') is True for c in all_cases),
        'human_verified_cases': 0, 'machine_matched_spans': machine_checked,
        'visually_verified_unmatched_spans': visual_checked, 'cross_split_near_duplicate_questions': 0,
        'status': 'agent_source_reviewed_human_pending', 'preceding_metadata_archive': str(archive.relative_to(output))}
    save_report(output / 'summary.json', summary)
    manifest = {**old_manifest, **summary, 'draft_revision': old_manifest.get('draft_revision', 0) + 1,
        'datasets': dataset_hashes, 'source_coverage': dict(coverage),
        'source_checks_sha256': sha((output / 'source-checks.json').read_bytes()),
        'summary_sha256': sha((output / 'summary.json').read_bytes()),
        'independent_review_artifacts': {name: sha((output / name).read_bytes())
            for name in ('dev-gold-audit.json', 'blind-gold-audit.json') if (output / name).exists()},
        'leakage_check': {'exact_span_overlap': 0, 'cluster_overlap': 0, 'template_group_overlap': 0,
                          'question_similarity_threshold': .92, 'near_duplicate_pairs': 0,
                          'cross_split_semantic_equivalence_review': 'automated checks do not prove semantic independence'},
        'review_status_note': 'Independent agents checked original sources and rewrote questions; no human-reviewed claim is made.'}
    save_report(output / 'manifest.json', manifest)
    (output / 'README.md').write_text('# Quality-v2 评测题集\n\n'
        '240题，开发/保留各120；跨来源、多轮、不可回答各24题；覆盖全部38份原资料。'
        '两位智能体分别独立核对原文件和改写各自题组，未使用RAG输出评分或调参。\n\n'
        '**agent_source_reviewed 不等于 human_verified。** 所有题目仍保留人工待核状态。'
        '纯视觉题不伪造文本span，文本指标仅使用有真实原文span的题。'
        '已检查精确span、证据簇、模板组跨split互斥及高相似问法；自动检查不证明语义完全独立。\n\n'
        'manifest.json冻结题集、source-checks.json和summary.json哈希；review-refresh-history保留重冻结前材料。'
        '开发集可调参；保留集只能在配置冻结后运行。\n', encoding='utf-8')
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command', nargs='?', choices=('build', 'refresh'), default='build')
    parser.add_argument('--root', type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument('--output', type=Path)
    args = parser.parse_args()
    operation = refresh if args.command == 'refresh' else build
    result = operation(args.root, args.output or args.root / '.rag/evaluation/quality-v2')
    print(json.dumps({k: result[k] for k in ('status', 'counts', 'covered_documents', 'questions_without_filename')}, ensure_ascii=False))


if __name__ == '__main__':
    main()
